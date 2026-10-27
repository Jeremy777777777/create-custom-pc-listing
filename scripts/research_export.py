"""Project scoped evidence into a workbook or check a reviewed export; never publish."""
import argparse
import hashlib
import json
import math
import tempfile
from pathlib import Path

import openpyxl
from jsonschema import Draft202012Validator
from research_values import check_fcc_id_value

ROOT = Path(__file__).resolve().parents[1]
FCC_KEY = ('Safety&Compliance', 'FCC ID')


def check_fcc_lookup(fact, require_outcome=False):
    """Check guide outcome semantics; source authentication remains human research."""
    evidence = fact.get('evidence', {})
    outcome = evidence.get('lookup_outcome')
    if outcome is None:  # Existing ledgers retain their existing validation contract.
        if require_outcome:
            raise ValueError('FCC export requires reconciled lookup_outcome evidence')
        return
    allowed = {'UNIT_CONFIRMED', 'MODEL_MODULE_MATCH', 'SUPPORTED_OPTION_SELECTED',
               'PENDING', 'FACTORY_NO_WLAN', 'HOST_SDOC'}
    if outcome not in allowed:
        raise ValueError('Unknown FCC lookup_outcome')
    if outcome in {'MODEL_MODULE_MATCH', 'SUPPORTED_OPTION_SELECTED', 'PENDING'} and fact['status'] == 'VERIFIED':
        raise ValueError('FCC model/support/pending outcome cannot confirm the sold unit')
    if outcome == 'FACTORY_NO_WLAN' and fact['status'] == 'VERIFIED':
        raise ValueError('Factory no-WLAN evidence is not an FCC identifier')
    if outcome in {'FACTORY_NO_WLAN', 'HOST_SDOC'} and fact.get('recommended_candidate') is not None:
        raise ValueError('No-radio/host-SDoC disposition cannot select a radio ID')
    if outcome == 'SUPPORTED_OPTION_SELECTED' and (fact['status'] != 'INPUT_UNVERIFIED'
                                                   or fact.get('recommended_candidate') is None):
        raise ValueError('Selected supported FCC option must remain an unverified candidate')
    if fact['status'] == 'NOT_APPLICABLE' and outcome not in {'FACTORY_NO_WLAN', 'HOST_SDOC'}:
        raise ValueError('FCC lookup outcome does not establish no applicable radio')
    if outcome == 'HOST_SDOC':
        if fact['status'] == 'VERIFIED' or fact.get('evidence_scope') != 'HOST' or fact.get('authorization_route') != 'SDoC':
            raise ValueError('Host SDoC cannot establish a certified radio FCC ID')
    if fact['status'] == 'NOT_APPLICABLE':
        if (fact.get('applies_to_sold_configuration') is not True
                or not fact.get('transfer_basis', '').strip()
                or not evidence.get('installation_basis') or not evidence.get('replacement_check')):
            raise ValueError('No-WLAN disposition needs sold applicability and later-addition check')
        if evidence.get('radio_inventory_checked') is not True or evidence.get('installed_radio_modules') != []:
            raise ValueError('FCC not-applicable disposition needs a checked empty installed radio inventory')
        if outcome == 'HOST_SDOC' and not evidence.get('sdoc_source'):
            raise ValueError('Host SDoC disposition needs the matched statement source')
        if outcome == 'FACTORY_NO_WLAN' and fact.get('evidence_scope') != 'RADIO_MODULE':
            raise ValueError('No-WLAN disposition must identify RADIO_MODULE scope')
        if evidence.get('missing_links'):
            raise ValueError('FCC not-applicable disposition cannot retain missing links')
    if outcome == 'UNIT_CONFIRMED' and fact['status'] == 'VERIFIED':
        from datetime import date
        from urllib.parse import urlparse
        model_key = 'host_model' if fact.get('evidence_scope') == 'HOST' else 'module_model'
        for key in (model_key, 'installation_basis', 'replacement_check', 'fcc_official_url', 'reviewed_date',
                    'fcc_record_id', 'fcc_record_model', 'fcc_applicant', 'fcc_grant_date'):
            if not isinstance(evidence.get(key), str) or not evidence[key].strip():
                raise ValueError('Confirmed FCC lookup requires evidence: ' + key)
        url = urlparse(evidence['fcc_official_url'])
        if url.scheme != 'https' or not (url.hostname == 'fcc.gov' or (url.hostname or '').endswith('.fcc.gov')):
            raise ValueError('FCC grant reference must use an official HTTPS FCC URL')
        reviewed = date.fromisoformat(evidence['reviewed_date'])
        granted = date.fromisoformat(evidence['fcc_grant_date'])
        if reviewed.isoformat() != evidence['reviewed_date'] or granted.isoformat() != evidence['fcc_grant_date']:
            raise ValueError('FCC dates must use YYYY-MM-DD')
        if granted > reviewed:
            raise ValueError('FCC grant date cannot follow its review date')
        if evidence.get('fcc_access_result') != 'REVIEWED':
            raise ValueError('Confirmed FCC ID needs a reviewed official record, not a failed lookup')
        if evidence['fcc_record_id'] != fact['final_value'] or evidence['fcc_record_model'] != evidence[model_key]:
            raise ValueError('FCC record ID/model must match the selected identifier and host/module')
        if evidence.get('missing_links'):
            raise ValueError('Confirmed FCC lookup cannot retain missing links')
# These existing template fields describe capability rather than installed capacity.
# All unlisted fields require SOLD_BASE evidence. Offer 'maximum price' is not hardware capability.
CAPABILITY_FIELDS = {
    ('Product Details', name) for name in (
        'Ram Memory Maximum Size', 'RAM Memory Maximum Size', 'RAM Memory Maximum Size Unit',
        'Display Maximum Resolution', 'Display Resolution Maximum', 'Display Resolution Maximum Unit',
        'Supported Monitor Maximum Quantity', 'CPU Model Speed Maximum', 'CPU Model Speed Maximum Unit',
        'Maximum Display Brightness')
}


def load_ledger(path):
    data = json.loads(Path(path).read_text(encoding='utf-8-sig'))
    validate_ledger(data)
    return data


def validate_ledger(data):
    schema = json.loads((ROOT/'assets/research-ledger-schema.json').read_text(encoding='utf-8'))
    errors = list(Draft202012Validator(schema).iter_errors(data))
    if errors:
        error = errors[0]
        raise ValueError('Scoped ledger schema: ' + '/'.join(map(str, error.absolute_path))
                         + ' violates ' + str(error.validator))
    from serial_research import validate_serial_research
    validate_serial_research(data)
    facts = {}
    for fact in data['facts']:
        if fact['id'] in facts:
            raise ValueError('Duplicate fact ID: ' + fact['id'])
        if fact['status'] not in ('VERIFIED', 'NOT_APPLICABLE') and fact['final_value'] is not None:
            raise ValueError('Unresolved fact must keep final_value null: ' + fact['id'])
        if fact['status'] == 'NOT_APPLICABLE' and fact['final_value'] is not None:
            raise ValueError('NOT_APPLICABLE must have a null value and source disposition')
        value = fact['final_value']
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('Research values must be finite: ' + fact['id'])
        if fact['status'] == 'VERIFIED' and (value is None or isinstance(value, str) and not value.strip()):
            raise ValueError('VERIFIED fact requires a value: ' + fact['id'])
        if isinstance(value, str) and value.lstrip().startswith('='):
            raise ValueError('Formula is not a research value: ' + fact['id'])
        has_candidate = 'value_candidate' in fact
        if has_candidate != ('value_candidate_source' in fact):
            raise ValueError('Candidate requires paired value_candidate_source: ' + fact['id'])
        if has_candidate:
            candidate = fact['value_candidate']
            if fact['status'] in ('VERIFIED', 'NOT_APPLICABLE') or value is not None:
                raise ValueError('Candidate must remain an unresolved fact: ' + fact['id'])
            if (candidate is None or isinstance(candidate, str) and
                    (not candidate.strip() or candidate.lstrip().startswith('='))
                    or isinstance(candidate, float) and not math.isfinite(candidate)):
                raise ValueError('Candidate must be a finite, nonempty literal: ' + fact['id'])
            for name in ('candidate_rationale', 'candidate_missing_confirmation'):
                text = fact.get('evidence', {}).get(name)
                if not isinstance(text, str) or not text.strip():
                    raise ValueError('Candidate requires evidence.' + name + ': ' + fact['id'])
            if fact.get('recommended_candidate') is not None:
                raise ValueError('Use one candidate representation per fact: ' + fact['id'])
        facts[fact['id']] = fact
    seen = set()
    for field in data['fields']:
        key = (field['sheet'], field['attribute'])
        if key in seen:
            raise ValueError('Duplicate field mapping: ' + str(key))
        seen.add(key)
        fact = facts.get(field['selected_fact'])
        if fact is None or (fact['sheet'], fact['attribute']) != key:
            raise ValueError('Missing or mismatched selected fact: ' + str(key))
        if fact['configurationId'] != data['configurationId']:
            raise ValueError('Selected fact belongs to another configuration: ' + str(key))
        allowed = {'SOLD_BASE', 'PLATFORM_CAPABILITY'} if key in CAPABILITY_FIELDS else {'SOLD_BASE'}
        if fact['fact_scope'] not in allowed:
            raise ValueError('Selected field requires sold/capability evidence, not factory/current-unit facts: ' + str(key))
        if fact['fact_scope'] == 'SOLD_BASE' and fact['status'] == 'VERIFIED':
            if fact.get('applies_to_sold_configuration') is not True or not fact.get('transfer_basis', '').strip():
                raise ValueError('Verified sold fact needs applicability and transfer_basis: ' + str(key))
        if key == ('Safety&Compliance', 'FCC ID') and fact['status'] == 'VERIFIED':
            check_fcc_id_value(fact['final_value'], 'VERIFIED')
            if fact.get('evidence_scope') not in ('HOST', 'RADIO_MODULE') or fact.get('authorization_route') != 'certification':
                raise ValueError('Verified FCC ID needs matched host/module certification scope')
        if key == FCC_KEY:
            check_fcc_lookup(fact)
    from research_audit import validate as audit_research
    audit_research(data)
    return facts


def project(data, field_keys, mode='draft', separate_candidates=True):
    if mode not in ('draft', 'export'):
        raise ValueError('Unknown projection mode')
    facts = validate_ledger(data)
    expected = list(field_keys)
    mapped = {(f['sheet'], f['attribute']): facts[f['selected_fact']] for f in data['fields']}
    if len(expected) != len(set(expected)) or set(mapped) != set(expected):
        raise ValueError('Ledger field mappings must cover the exact current template')
    projected = {}
    for key in expected:
        fact = mapped[key]
        candidate = fact.get('recommended_candidate')
        advisory = fact.get('value_candidate')
        advisory_source = None
        if advisory is not None:
            if key == FCC_KEY:
                check_fcc_id_value(advisory, 'VERIFIED')
            if not separate_candidates:
                raise ValueError('Candidate metadata requires the seven-row workbook layout')
            evidence = fact['evidence']
            advisory_source = ('MOST_LIKELY_CANDIDATE — NOT CONFIRMED; ' + fact['value_candidate_source']
                               + '; Rationale: ' + evidence['candidate_rationale']
                               + '; Missing confirmation: ' + evidence['candidate_missing_confirmation'])
        if mode == 'export':
            if fact['status'] not in ('VERIFIED', 'NOT_APPLICABLE'):
                raise ValueError('Export blocked by unresolved field: ' + str(key))
            if candidate is not None or advisory is not None:
                raise ValueError('Export blocked by recommended candidate: ' + str(key))
            if key == FCC_KEY:
                check_fcc_lookup(fact, require_outcome=True)
        value = fact['final_value']
        source = fact['source']
        if mode == 'draft' and candidate is not None:
            if fact['status'] != 'INPUT_UNVERIFIED' or key != ('Safety&Compliance', 'FCC ID'):
                raise ValueError('Only INPUT_UNVERIFIED FCC research cells may show a recommended candidate')
            if not isinstance(candidate, str) or not candidate.strip() or candidate.lstrip().startswith('='):
                raise ValueError('FCC candidate must be a nonempty text identifier')
            check_fcc_id_value(candidate, 'VERIFIED')
            warning = 'MOST_LIKELY_CANDIDATE — NOT CONFIRMED; '
            if fact.get('evidence', {}).get('lookup_outcome') == 'SUPPORTED_OPTION_SELECTED':
                warning = 'USER_SELECTED_SUPPORTED_OPTION — INSTALLATION NOT CONFIRMED; '
            if separate_candidates:
                advisory, advisory_source = candidate, warning + source
            else:  # Check preserved five-row workbooks without rewriting their history.
                value, source = candidate, warning + source
        if key == FCC_KEY and fact.get('evidence', {}).get('lookup_outcome'):
            evidence = fact['evidence']
            details = ['FCC_LOOKUP=' + evidence['lookup_outcome']]
            details += [name + '=' + str(evidence[name]) for name in
                        ('module_model', 'host_model', 'fcc_record_id', 'fcc_record_model',
                         'fcc_applicant', 'fcc_grant_date', 'fcc_access_result',
                         'fcc_official_url', 'reviewed_date') if evidence.get(name)]
            source += '; ' + '; '.join(details)
        projected[key] = {'value': value, 'status': fact['status'], 'source': source,
                          'value_candidate': advisory, 'value_candidate_source': advisory_source,
                          'fact_id': fact['id'], 'fact_scope': fact['fact_scope']}
    return projected


def workbook_fields(book):
    return [(sheet.title, cell.value) for sheet in book for cell in list(sheet[1])[1:]]


def same_cell_value(actual, expected):
    """Allow Excel's int/float round trip without conflating Boolean and number."""
    if isinstance(expected, bool) or isinstance(actual, bool):
        return type(actual) is type(expected) and actual == expected
    if isinstance(expected, (int, float)):
        return isinstance(actual, (int, float)) and math.isfinite(actual) and actual == expected
    return type(actual) is type(expected) and actual == expected


def verify_workbook(book, data, mode='draft'):
    layouts = [tuple(sheet.cell(r, 1).value for r in (6, 7)) for sheet in book]
    modern = all(labels == ('Value Candidate', 'Value Candidate Source') for labels in layouts)
    legacy = all(labels == (None, None) for labels in layouts)
    if not modern and not legacy:
        raise ValueError('Workbook candidate row labels are incomplete or changed')
    if mode == 'export' and not modern:
        raise ValueError('Export requires the current seven-row workbook layout')
    result = project(data, workbook_fields(book), mode, separate_candidates=modern)
    rows = ((3, 'value'), (4, 'status'), (5, 'source'))
    if modern:
        rows += ((6, 'value_candidate'), (7, 'value_candidate_source'))
    for sheet in book:
        for column in range(2, sheet.max_column + 1):
            key = (sheet.title, sheet.cell(1, column).value)
            field = result[key]
            if not all(same_cell_value(sheet.cell(row, column).value, field[name])
                       for row, name in rows):
                raise ValueError('Workbook differs from scoped ledger projection: ' + str(key))
            if any(sheet.cell(row, column).data_type == 'f' for row, _ in rows):
                raise ValueError('Formula cannot be exported: ' + str(key))
    return result


def write_workbook(template, ledger, workbook, mode='draft', product_id=None,
                   configuration_id=None, stage='RESEARCH_REVIEW', previous_workbook=None):
    """Create a new canonical workbook; validate before installing the output."""
    template, ledger, workbook = map(Path, (template, ledger, workbook))
    from workbook_paths import files_with_suffix
    if workbook.suffix.casefold() != '.xlsx' or not workbook.parent.is_dir():
        raise ValueError('Output requires an existing product directory and .xlsx suffix')
    if any(p.name.casefold() == workbook.name.casefold() for p in files_with_suffix(workbook.parent, '.xlsx')):
        raise ValueError('Output already exists; use a new staged filename')
    if any(p.casefold() == 'workbook-snapshots' for p in workbook.resolve().parts):
        raise ValueError('Cannot write historical workbook snapshots')
    data = load_ledger(ledger)
    from serial_research import validate_serial_research
    validate_serial_research(data, required=True, full_research=True)
    from research_handoff import check_product_location
    check_product_location(workbook, data['productId'])
    book = openpyxl.load_workbook(template)
    if any(tuple(sheet.cell(r, 1).value for r in (6, 7)) !=
           ('Value Candidate', 'Value Candidate Source') for sheet in book):
        book.close()
        raise ValueError('New workbooks require the seven-row candidate template')
    previous = None
    temporary = None
    try:
        projected = project(data, workbook_fields(book), mode)
        from dimension_units import validate_dimensions
        validate_dimensions(projected)
        selected = validate_ledger(data)
        fcc = next((selected[f['selected_fact']] for f in data['fields']
                    if (f['sheet'], f['attribute']) == FCC_KEY), None)
        if fcc and not fcc.get('evidence', {}).get('lookup_outcome'):
            raise ValueError('New FCC workbook requires guide lookup_outcome metadata')
        if workbook.parent.parent.name == 'product generated photo':
            from research_handoff import check_research_summaries
            check_research_summaries(workbook.parent.parent.parent, product_id=data['productId'])
            # All top-level product workbooks are current, including earlier stages.
            # A revised ledger must not leave an old stage silently stale until CI.
            for sibling in files_with_suffix(workbook.parent, '.xlsx'):
                old = openpyxl.load_workbook(sibling)
                try:
                    verify_workbook(old, data)
                except ValueError as error:
                    raise ValueError('Stale current workbook; preserve/reconcile it before generating: '
                                     + sibling.name) from error
                finally:
                    old.close()
        if previous_workbook:
            from workbook_paths import resolve_workbook_path
            previous_path = resolve_workbook_path(previous_workbook)
            if previous_path.parent.resolve() != workbook.parent.resolve():
                bundle = previous_path.parent.resolve()
                if bundle.parent != workbook.parent.resolve()/'workbook-snapshots':
                    raise ValueError('Previous workbook must belong to the same product directory or its snapshot')
                from workbook_snapshots import check_snapshot
                binding = check_snapshot(bundle)
                receipt = json.loads((bundle/'snapshot.json').read_text(encoding='utf-8-sig'))
                if (any(binding[key] != data[key] for key in ('productId', 'configurationId'))
                        or receipt['files']['workbook']['path'] != previous_path.name):
                    raise ValueError('Previous snapshot workbook/product/configuration does not match')
            previous = openpyxl.load_workbook(previous_path)
            if workbook_fields(previous) != workbook_fields(book):
                raise ValueError('Previous workbook must use the same template fields')
        from copy import copy
        from openpyxl.styles import PatternFill
        for sheet in book:
            for col in range(2, sheet.max_column + 1):
                key = (sheet.title, sheet.cell(1, col).value)
                field = projected[key]
                for row, name in ((3, 'value'), (4, 'status'), (5, 'source'),
                                  (6, 'value_candidate'), (7, 'value_candidate_source')):
                    cell = sheet.cell(row, col)
                    cell.value = field[name]
                    # Evidence text beginning '=' must stay text, never executable Excel formulas.
                    if isinstance(field[name], str):
                        cell.data_type = 's'
                if key == FCC_KEY:
                    sheet.cell(6, col).number_format = '@'
                    cell = sheet.cell(3, col)
                    cell.number_format = '@'
                    old = previous[sheet.title] if previous else None
                    historical = (old is not None and field['status'] == 'VERIFIED'
                                  and old.cell(4, col).value == 'VERIFIED'
                                  and old.cell(3, col).value == field['value'])
                    cell.fill = copy(old.cell(3, col).fill) if historical else PatternFill(
                        'solid', fgColor='C6EFCE' if field['value'] else 'FFC7CE')
        from workbook_reference_audit import validate as validate_references, urls
        from workbook_reference_audit import check_paths as check_reference_paths
        check_reference_paths(ledger, required=True)
        reference_audit = json.loads((ledger.parent/'reference-audit.json').read_text(encoding='utf-8-sig'))
        reference_result = validate_references(data, reference_audit, ledger.read_bytes())
        for sheet in book:
            for col in range(2, sheet.max_column + 1):
                key = (sheet.title, sheet.cell(1, col).value)
                for row in (5, 7):
                    cell = sheet.cell(row, col)
                    targets = urls(cell.value)
                    cell.hyperlink = (reference_result['links'][key] if row == 5 else targets[0]) if targets else None
        validate_references(data, reference_audit, ledger.read_bytes(), book)
        with tempfile.NamedTemporaryFile(prefix='.' + workbook.stem + '-', suffix='.xlsx',
                                         dir=workbook.parent, delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        result = run(template, ledger, temporary, mode, product_id, configuration_id, stage)
        # Exclusive creation protects previous deliverables even if a file appears during validation.
        payload = temporary.read_bytes()
        created = False
        try:
            with workbook.open('xb') as output:
                created = True
                output.write(payload)
        except BaseException:
            if created:
                workbook.unlink(missing_ok=True)
            raise
        return result
    finally:
        book.close()
        if previous:
            previous.close()
        if temporary:
            temporary.unlink(missing_ok=True)


def run(template, ledger, workbook, mode, product_id=None, configuration_id=None, stage=None, check_copy=False):
    template, ledger, workbook = map(Path, (template, ledger, workbook))
    from workbook_paths import resolve_workbook_path
    workbook=resolve_workbook_path(workbook)
    from research_history import check_product_research_history
    check_product_research_history(workbook.parent)
    canonical_template = ROOT/'assets/listing-workbook-template.xlsx'
    input_hashes = {name: hashlib.sha256(Path(path).read_bytes()).hexdigest()
                    for name, path in (('workbook', workbook), ('ledger', ledger), ('template', template))}
    # Import only at execution time; check-workbooks can use verify_workbook without a cycle.
    import importlib.util
    spec = importlib.util.spec_from_file_location('workbook_checks', ROOT/'scripts/check-workbooks.py')
    checks = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(checks)
    if mode == 'export':
        if stage != 'LISTING_DRAFT':
            raise ValueError('Export requires explicit LISTING_DRAFT stage')
        if any(part.casefold() == 'workbook-snapshots' for part in workbook.resolve().parts):
            raise ValueError('Historical workbook snapshots cannot be exported')
        if input_hashes['template'] != hashlib.sha256(canonical_template.read_bytes()).hexdigest():
            raise ValueError('Export requires the current canonical template; alternate/old templates are draft-only')
        checks.check_template_inputs(canonical_template)
    product = None
    metadata_report = None
    manufacturer = None
    data = load_ledger(ledger)
    if mode == 'export':
        from serial_research import validate_serial_research
        validate_serial_research(data, required=True, full_research=True)
    from research_handoff import check_product_location,coverage_summary,internal_storage_configuration
    check_product_location(workbook,data['productId'])
    listing_stage = stage == 'LISTING_DRAFT' if stage else '_LISTING_' in workbook.stem.upper()
    if check_copy and not listing_stage:
        raise ValueError('Copy checking requires LISTING_DRAFT stage')
    facts_path=workbook.parent/'product-facts.json'
    if facts_path.exists():
        input_hashes['productFacts']=hashlib.sha256(facts_path.read_bytes()).hexdigest()
        raw_product=json.loads(facts_path.read_text(encoding='utf-8-sig'))
        # Research may be incomplete, but an existing identity cannot disagree.
        from product_identity import product_id as metadata_product_id
        for value,expected in ((metadata_product_id(raw_product),data['productId']),
                               (raw_product.get('selectedBaseConfiguration',{}).get('id'),data['configurationId'])):
            if value is not None and value != expected:
                raise ValueError('Product facts identity/configuration differs from research target')
    if mode == 'export' or listing_stage:
        facts_path=workbook.parent/'product-facts.json'
        input_hashes['productFacts']=hashlib.sha256(facts_path.read_bytes()).hexdigest()
        from product_metadata import listing_view
        raw_product = json.loads(facts_path.read_text(encoding='utf-8-sig'))
        product,metadata_report = listing_view(raw_product,data)
        manufacturer = product['identity'].get('oemBrand')
        if not manufacturer:
            raise ValueError('Listing workbook requires OEM brand for existing copy checks')
    checks.check_output(workbook, checks.template_topology(template), manufacturer)
    if mode == 'export' and (not product_id or not configuration_id):
        raise ValueError('Export requires explicit target product ID and configuration ID')
    if product_id is not None and product_id != data['productId']:
        raise ValueError('Ledger belongs to another target product')
    if configuration_id is not None and configuration_id != data['configurationId']:
        raise ValueError('Ledger belongs to another target configuration')
    if mode == 'export':
        if product.get('identity',{}).get('internalId')!=data['productId'] or product.get('selectedBaseConfiguration',{}).get('id')!=data['configurationId']:
            raise ValueError('Product facts identity/configuration differs from export target')
    book = openpyxl.load_workbook(workbook, data_only=False)
    try:
        projected = verify_workbook(book, data, mode)
        if mode == 'export' or check_copy:
            from listing_content import validate_listing
            from capacity_units import STORAGE_FIELDS,check_nominal_storage_basis
            selected={f['id']:f for f in data['facts']}
            for mapping in data['fields']:
                if mapping['sheet']=='Product Details' and mapping['attribute'] in STORAGE_FIELDS:
                    check_nominal_storage_basis(selected[mapping['selected_fact']])
            definitions={(sheet.title,sheet.cell(1,c).value):sheet.cell(2,c).value
                         for sheet in book for c in range(2,sheet.max_column+1)}
            storage=internal_storage_configuration(data)
            validate_listing(projected,product,definitions,storage_total_gb=storage['totalGB'] if storage else None,storage_configuration=storage)
    finally:
        book.close()
    for name, path in (('workbook', workbook), ('ledger', ledger), ('template', template)):
        if hashlib.sha256(Path(path).read_bytes()).hexdigest() != input_hashes[name]:
            raise ValueError('Input changed during validation: ' + name)
    if 'productFacts' in input_hashes and hashlib.sha256(facts_path.read_bytes()).hexdigest()!=input_hashes['productFacts']:
        raise ValueError('Input changed during validation: productFacts')
    from research_audit import validate as audit_research
    from workbook_reference_audit import check_paths as check_references
    reference_result = check_references(ledger, workbook, required=mode == 'export')
    return {'referenceAudit': reference_result, 'researchTrace': audit_research(data),
            'productFactsSha256':input_hashes.get('productFacts'), 'result': 'EXPORT_EVIDENCE_CHECK_PASS' if mode == 'export' else 'DRAFT_PROJECTION_PASS',
            'productId': data['productId'], 'configurationId': data['configurationId'],
            'workbookSha256': input_hashes['workbook'],
            'ledgerSha256': input_hashes['ledger'],
            'templateSha256': input_hashes['template'],
            'stage': stage or ('LISTING_DRAFT' if listing_stage else 'RESEARCH_REVIEW'),
            'amazonEligibilityAssessed': False, 'metadataView': metadata_report,
            'coverageSummary':coverage_summary(data,projected),
            'fields': [{'sheet': k[0], 'attribute': k[1], **v} for k,v in projected.items()]}


def summary_destination(output, workbook, identity):
    """Preflight the sidecar binding before creating any new deliverables."""
    from research_handoff import check_product_location
    from workbook_paths import files_with_suffix
    output, workbook = Path(output), Path(workbook)
    if output.parent.resolve() != workbook.parent.resolve():
        raise ValueError('Research summary must be beside its workbook')
    check_product_location(output, identity['productId'])
    if not output.name.casefold().endswith('.research-summary.json'):
        raise ValueError('Summary output must end with .research-summary.json')
    matches = [p for p in files_with_suffix(output.parent, '.research-summary.json')
               if p.name.casefold() == output.name.casefold()]
    if matches:
        output = matches[0]
    if output.exists():
        previous = json.loads(output.read_text(encoding='utf-8-sig'))
        binding = previous.get('workbookFile') if isinstance(previous, dict) else None
        if (isinstance(previous, dict) and binding is None and workbook.is_file()
                and previous.get('workbookSha256') == hashlib.sha256(workbook.read_bytes()).hexdigest()):
            binding = workbook.name  # Safely upgrade an older hash-bound sidecar.
        if (not isinstance(previous, dict) or previous.get('artifactRole') != 'RESEARCH_COVERAGE_SUMMARY'
                or any(previous.get(key) != identity[key] for key in ('productId', 'configurationId'))
                or str(binding or '').casefold() != workbook.name.casefold()):
            raise ValueError('Summary output would overwrite an unrelated workbook/configuration artifact')
    return output


def write_summary(output, result, workbook, ledger=None):
    """Replace the sidecar only after fully preparing its bytes."""
    serial = result.get('coverageSummary', {}).get('serialResearch', {})
    if (not serial.get('recordPresent') or serial.get('outcome') == 'NOT_ASSESSED_LEGACY'
            or serial.get('workflow_mode') != 'FULL_LISTING_WORKFLOW'):
        raise ValueError('New workbook summary requires assessed FULL_LISTING_WORKFLOW serial_research')
    from workbook_paths import resolve_workbook_path
    summary = {k: result[k] for k in ('productId', 'configurationId', 'stage', 'workbookSha256',
               'ledgerSha256', 'templateSha256', 'productFactsSha256', 'coverageSummary')}
    summary['artifactRole'] = 'RESEARCH_COVERAGE_SUMMARY'
    summary['productFactsBindingVersion'] = 1
    summary['workbookFile'] = resolve_workbook_path(workbook).name
    if 'researchTrace' in result:
        summary['researchTrace'] = result['researchTrace']
        summary['researchTraceSchemaVersion'] = 2
        if ledger is not None:
            from research_audit import validate_summary_trace
            validate_summary_trace(load_ledger(ledger), summary['researchTrace'])
    if output.exists():
        previous = json.loads(output.read_text(encoding='utf-8-sig'))
        old_trace = previous.get('researchTrace', {})
        if not isinstance(old_trace, dict):
            raise ValueError('Invalid prior research trace summary')
        if old_trace.get('legacyReuse') and not summary.get('researchTrace', {}).get('legacyReuse'):
            bound = ('productId', 'configurationId', 'stage', 'workbookSha256', 'ledgerSha256',
                     'templateSha256', 'productFactsSha256')
            if any(previous.get(key) != summary.get(key) for key in bound):
                raise ValueError('Retained research inputs changed; supply previous ledger and reuse review to refresh')
            from research_audit import validate_summary_trace
            if ledger is None:
                raise ValueError('Ledger required to preserve a retained research summary')
            validate_summary_trace(load_ledger(ledger), old_trace)
            summary['researchTrace'] = old_trace
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix='.' + output.name + '-', suffix='.tmp',
                                         dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write((json.dumps(summary, indent=2, ensure_ascii=False) + '\n').encode('utf-8'))
        temporary.replace(output)
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)
    return summary


def main():
    import sys
    # Keep redirected JSON deterministic on Windows as well as Linux.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path, default=ROOT/'assets/listing-workbook-template.xlsx')
    parser.add_argument('--ledger', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--mode', choices=['draft', 'export'], default='draft')
    parser.add_argument('--product-id', help='Required for export: independently selected target')
    parser.add_argument('--configuration-id', help='Required for export: target base configuration')
    parser.add_argument('--stage', choices=['RESEARCH_REVIEW', 'LISTING_DRAFT'],
                        help='Export requires explicit LISTING_DRAFT; filenames cannot grant export eligibility')
    parser.add_argument('--summary-output',type=Path,help='Write a standalone research coverage JSON after successful checking')
    parser.add_argument('--write-workbook', action='store_true', help='Create a new Excel from the scoped ledger, then validate it')
    parser.add_argument('--previous-workbook', type=Path, help='Preserve FCC coloring from a current workbook or checked same-product snapshot')
    parser.add_argument('--previous-ledger', type=Path, help='Observed previous scoped ledger for a bounded legacy revision')
    parser.add_argument('--research-reuse', type=Path, help='Hash-bound applicability review for unchanged legacy fields')
    args = parser.parse_args()
    created_workbook = False
    try:
        if args.previous_workbook and not args.write_workbook:
            raise ValueError('--previous-workbook requires --write-workbook')
        if bool(args.previous_ledger) != bool(args.research_reuse):
            raise ValueError('--previous-ledger and --research-reuse require each other')
        if args.write_workbook:
            from research_audit import validate as audit_research
            trace_result = audit_research(load_ledger(args.ledger), required=True,
                previous_bytes=args.previous_ledger.read_bytes() if args.previous_ledger else None,
                reuse=json.loads(args.research_reuse.read_text(encoding='utf-8-sig')) if args.research_reuse else None)
            if not args.summary_output:
                args.summary_output = args.workbook.with_suffix('.research-summary.json')
            args.summary_output = summary_destination(args.summary_output, args.workbook, load_ledger(args.ledger))
            result = write_workbook(args.template, args.ledger, args.workbook, args.mode,
                                    args.product_id, args.configuration_id,
                                    args.stage or 'RESEARCH_REVIEW', args.previous_workbook)
            created_workbook = True
            result['researchTrace'] = trace_result
        else:
            result = run(args.template, args.ledger, args.workbook, args.mode, args.product_id, args.configuration_id, args.stage)
            if args.research_reuse:
                from research_audit import validate as audit_research
                result['researchTrace'] = audit_research(load_ledger(args.ledger), required=True,
                    previous_bytes=args.previous_ledger.read_bytes(),
                    reuse=json.loads(args.research_reuse.read_text(encoding='utf-8-sig')))
        if args.summary_output:
            output = summary_destination(args.summary_output, args.workbook, result)
            if output.exists() and not args.research_reuse:
                from research_audit import validate_summary_trace
                prior = json.loads(output.read_text(encoding='utf-8-sig'))
                if 'researchTrace' in prior and prior.get('ledgerSha256') == result['ledgerSha256']:
                    validate_summary_trace(load_ledger(args.ledger), prior['researchTrace'])
                elif 'researchTrace' not in prior and prior.get('researchTraceSchemaVersion') is not None:
                    raise ValueError('Research trace missing from a versioned summary')
            saved = write_summary(output, result, args.workbook, args.ledger)
            result['researchTrace'] = saved['researchTrace']
    except (ValueError, OSError, json.JSONDecodeError) as error:
        if created_workbook:
            try:
                args.workbook.unlink(missing_ok=True)
            except OSError as cleanup_error:
                parser.exit(1, str(error) + '; workbook cleanup failed at ' + str(args.workbook)
                            + ': ' + str(cleanup_error) + '\n')
        parser.exit(1, str(error) + '\n')
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
