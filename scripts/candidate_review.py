"""Visible review values, independent of confirmed Amazon numeric projection."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = 'Candidate Review'
NOTICE = ('Research candidates only — NOT CONFIRMED. Original units are retained. '
          'Use the field links for evidence and alternatives; these are not approved Amazon entries.')


def data_sheets(book):
    return [s for s in book if s.title != SHEET]


def cell_value(record, review=False):
    return record.get('candidate_display', record.get('value_candidate')) if review else record.get('value_candidate')


def rows(data, projected):
    from candidate_ranking import assess, display
    facts = {f['id']: f for f in data['facts']}
    result = []
    for key, record in projected.items():
        if record['status'] in {'VERIFIED', 'NOT_APPLICABLE'}:
            continue
        fact = facts[record['fact_id']]
        from fcc_inventory import is_external_field
        if is_external_field(data, fact):
            continue
        ranked = assess(fact)
        if record.get('hardware_reference'):
            value = display(record['hardware_reference'])
            meaning, priority = 'Recorded hardware reference; AMAZON ENTRY UNCONFIRMED; ' + str(record['hardware_reference'].get('scope')), 0
        elif ranked:
            options = ranked['options']
            primary = next((o for o in options if o['id'] == ranked.get('primary_id')), None)
            value = display(primary) if primary else ' / '.join(display(o) for o in options)
            basis = primary.get('basis', fact.get('evidence', {}).get('candidate_basis')) if primary else 'TIED'
            meaning = ('QUALIFIED ' + basis + ' REFERENCE' if basis in {'OEM_MAXIMUM', 'FACTORY_ONLY', 'FAMILY_CAPABILITY'}
                       else (primary['confidence'] + ' confidence' if primary else 'TIED — no supported single winner'))
            meaning += '; NOT CONFIRMED'
            if record.get('value_candidate') is None:
                meaning += '; Amazon numeric entry withheld'
            priority = 0 if record.get('value_candidate') is None else 1
        elif record.get('value_candidate') is not None:
            value, meaning, priority = str(record['value_candidate']), 'NOT CONFIRMED — see source and qualification', 1
        else:
            review = fact.get('evidence', {}).get('candidate_review', {})
            value = review.get('decision', 'RESEARCH_PENDING')
            meaning, priority = review.get('missing_confirmation', 'See field evidence and next action'), 2
        result.append((priority, key, (key[1], value, meaning, 'Open field and sources')))
    return sorted(result, key=lambda row: row[0])


def expected(data, projected, book):
    result = {(1, 1): data['productId'] + ' — Candidate Review', (2, 1): NOTICE}
    for col, label in enumerate(('Field', 'Recommended candidate / remaining input', 'Qualification', 'Evidence'), 1):
        result[(4, col)] = label
    links = {}
    for row, (_, key, values) in enumerate(rows(data, projected), 5):
        for col, value in enumerate(values, 1):
            result[(row, col)] = value
        source = book[key[0]]
        column = next(c.column for c in source[1] if c.value == key[1])
        links[(row, 4)] = "#'" + key[0].replace("'", "''") + "'!" + get_column_letter(column) + '6'
    return result, links


def write(book, data, projected):
    if SHEET in book.sheetnames:
        raise ValueError('Do not replace an existing candidate review silently')
    sheet = book.create_sheet(SHEET, 0)
    values, links = expected(data, projected, book)
    for (row, col), value in values.items():
        cell = sheet.cell(row, col, value)
        cell.data_type = 's'
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cell.font = Font(name='Calibri', size=12, color='172C40')
        if row == 1 or row == 4:
            cell.fill = PatternFill('solid', fgColor='183B56')
            cell.font = Font(name='Calibri', size=15 if row == 1 else 12, bold=True, color='FFFFFF')
        elif row >= 5:
            cell.fill = PatternFill('solid', fgColor='FFF2CC' if row % 2 else 'FFF8E7')
    sheet.merge_cells('A1:D1'); sheet.merge_cells('A2:D2')
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 48
    sheet.row_dimensions[4].height = 32
    for row in range(5, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 76
    for col, width in zip('ABCD', (30, 48, 50, 23)):
        sheet.column_dimensions[col].width = width
    for position, target in links.items():
        sheet.cell(*position).hyperlink = target
        sheet.cell(*position).font = Font(name='Calibri', size=12, color='0563C1', underline='single')
    sheet.freeze_panes = 'B5'
    sheet.auto_filter.ref = 'A4:D' + str(sheet.max_row)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    for other in book:
        other.sheet_view.tabSelected = other is sheet
    book.active = 0
    book.views[0].firstSheet = 0


def verify(book, data, projected, required=False):
    if SHEET not in book.sheetnames:
        if required:
            raise ValueError('New draft requires an opening Candidate Review with visible recommendations')
        return
    sheet = book[SHEET]
    values, links = expected(data, projected, book)
    if book.sheetnames[0] != SHEET or book.active.title != SHEET or sheet.sheet_state != 'visible':
        raise ValueError('Candidate Review must be the visible opening sheet')
    actual = {(c.row, c.column): c.value for row in sheet for c in row if c.value is not None}
    if actual != values:
        raise ValueError('Candidate Review differs from scoped candidate projection')
    for row in sheet:
        for cell in row:
            target = cell.hyperlink.target if cell.hyperlink else None
            if cell.data_type == 'f' or target != links.get((cell.row, cell.column)):
                raise ValueError('Candidate Review contains a formula or changed evidence link')
    if sheet.sheet_view.topLeftCell not in (None, 'A1'):
        raise ValueError('Opening candidate recommendations must be immediately visible')
    if (sheet.row_dimensions[5].hidden or sheet.column_dimensions['B'].hidden
            or book.views[0].firstSheet not in (None, 0)):
        raise ValueError('Opening candidate row/column cannot be hidden')
