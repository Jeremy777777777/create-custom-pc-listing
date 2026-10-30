"""Check workbook structure and obvious FCC ID placeholders; never authenticate compliance."""
import argparse
import json
import os
from pathlib import Path
import openpyxl
from warranty_policy import validate_workbook_copy
from research_values import check_fcc_id_value
from workbook_snapshots import check_snapshots, check_snapshot_history

SHEETS = ['Product Details', 'Offer', 'Safety&Compliance']
RESEARCH_STATUSES = ('VERIFIED', 'INPUT_UNVERIFIED', 'CONFLICT', 'TBD',
                     'SOURCE_UNAVAILABLE', 'NOT_APPLICABLE')
FCC_FIELDS = ('FCC Radio Frequency Emission Compliance',
              'Radio Frequency Emission & Authorization Status', 'FCC ID',
              'SDoC Contact Name', 'SDoC Contact US Mailing Address',
              'SDoC Contact Email Address', 'SDOC Contact US Phone Number')


def check_template_inputs(path):
    """Check required FCC inputs and the shared research status vocabulary."""
    book = openpyxl.load_workbook(path)
    try:
        sheet = book['Safety&Compliance']
        headers = [cell.value for cell in sheet[1]]
        if any(headers.count(field) != 1 for field in FCC_FIELDS):
            raise ValueError('Template requires one of each FCC/SDoC field')
        col = headers.index('FCC ID') + 1
        if sheet.cell(3, col).number_format != '@':
            raise ValueError('Template FCC ID Value must use text formatting')
        for tab in book:
            if tuple(tab.cell(r, 1).value for r in (6, 7)) != ('Value Candidate', 'Value Candidate Source'):
                raise ValueError('Template requires Value Candidate and Value Candidate Source rows')
            for column in range(2, tab.max_column + 1):
                cell = tab.cell(4, column)
                if not any(cell.coordinate in validation and validation.type == 'list'
                           and validation.formula1 == '"' + ','.join(RESEARCH_STATUSES) + '"'
                           for validation in tab.data_validations.dataValidation):
                    raise ValueError('Template Status dropdown must match research statuses')
    finally:
        book.close()


def check_field_disposition(status, source):
    if status not in RESEARCH_STATUSES:
        raise ValueError('Status must be one of the six research statuses')
    if not isinstance(source, str) or not source.strip():
        raise ValueError('Source or disposition must contain non-whitespace text')

def template_topology(path):
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        if book.sheetnames != SHEETS:
            raise ValueError('Unexpected template sheets/order')
        result = {}
        for sheet in book:
            if sheet.max_row < 5 or sheet.max_column < 2:
                raise ValueError(f'{sheet.title}: expected five template rows')
            headers = [sheet.cell(1,c).value for c in range(2,sheet.max_column+1)]
            if any(not h for h in headers) or len(set(headers)) != len(headers):
                raise ValueError(f'{sheet.title}: missing or duplicate field names')
            result[sheet.title] = ([tuple(sheet.cell(r,c).value for c in range(1,sheet.max_column+1)) for r in (1,2)],
                                   tuple(sheet.cell(r,1).value for r in range(3,8)
                                         if r <= 5 or sheet.cell(r,1).value is not None))
        return result
    finally:
        book.close()

def check_output(path, topology, manufacturer=None, listing_mode=None, ledger_path=None, verify_projection=True):
    ledger_path = Path(ledger_path) if ledger_path else path.parent/'research-ledger.json'
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        source_only_projection_checked = False
        from candidate_review import data_sheets
        sheets = data_sheets(book)
        if len(sheets) != len(book.worksheets) and not ledger_path.is_file():
            raise ValueError('Candidate Review requires its scoped research ledger')
        if [s.title for s in sheets] != list(topology):
            raise ValueError(f'{path}: sheets/order changed')
        for sheet in sheets:
            rows, labels = topology[sheet.title]
            if sheet.max_row < 5 or sheet.max_column != len(rows[0]):
                raise ValueError(f'{path}: {sheet.title} template dimensions changed')
            if [tuple(sheet.cell(r,c).value for c in range(1,sheet.max_column+1)) for r in (1,2)] != rows:
                raise ValueError(f'{path}: {sheet.title} headers/definitions changed')
            modern = tuple(sheet.cell(r,1).value for r in (6,7)) == ('Value Candidate', 'Value Candidate Source')
            actual_labels = tuple(sheet.cell(r,1).value for r in range(3,8 if modern else 6))
            if actual_labels != (labels if modern else labels[:3]):
                raise ValueError(f'{path}: {sheet.title} row labels changed')
            if any(sheet.cell(r,c).value is not None for r in range(8 if modern else 6,sheet.max_row+1) for c in range(1,sheet.max_column+1)):
                raise ValueError(f'{path}: {sheet.title} extra data rows added')
            for c in range(2,sheet.max_column+1):
                if not sheet.cell(4,c).value or not sheet.cell(5,c).value:
                    raise ValueError(f'{path}: {sheet.title}/{sheet.cell(1,c).value} lacks Status/Source or disposition')
                check_field_disposition(sheet.cell(4,c).value, sheet.cell(5,c).value)
                if modern:
                    candidate, candidate_source = sheet.cell(6,c), sheet.cell(7,c)
                    if any(sheet.cell(r,c).data_type == 'f' for r in (3,4,5,6,7)):
                        raise ValueError('Formula cannot be used in workbook research rows')
                    if (candidate.value is None) != (candidate_source.value is None):
                        # A recorded no-candidate disposition or unit-qualified
                        # hardware recommendation may legitimately occupy row 7
                        # alone. Authenticate its exact projection from the scoped
                        # ledger; a free-form note/prefix is never an exemption.
                        if candidate.value is not None or not ledger_path.is_file():
                            raise ValueError('Candidate requires paired source or a scoped projected disposition')
                        if not source_only_projection_checked:
                            from research_export import load_ledger, verify_workbook
                            verify_workbook(book, load_ledger(ledger_path))
                            source_only_projection_checked = True
                    if candidate.value is not None:
                        if (sheet.cell(3,c).value is not None or sheet.cell(4,c).value in ('VERIFIED','NOT_APPLICABLE')
                                or not isinstance(candidate_source.value,str) or not candidate_source.value.strip()):
                            raise ValueError('Candidate must remain separate from confirmed Value')
                if sheet.title == 'Safety&Compliance' and sheet.cell(1,c).value == 'FCC ID':
                    check_fcc_id_value(sheet.cell(3,c).value, sheet.cell(4,c).value)
        if manufacturer is not None:
            facts_path = path.parent/'product-facts.json'
            if facts_path.exists():
                from listing_mode import resolve
                observed_mode = resolve(json.loads(facts_path.read_text(encoding='utf-8-sig')))
                if listing_mode is not None and listing_mode != observed_mode:
                    raise ValueError('Copy mode differs from product facts')
                listing_mode = observed_mode
            validate_workbook_copy(book, manufacturer, listing_mode)
        if ledger_path.exists():
            ledger = json.loads(ledger_path.read_text(encoding='utf-8-sig'))
            if isinstance(ledger, dict):
                from research_export import verify_workbook
                from research_handoff import check_product_location
                facts_path=path.parent/'product-facts.json'
                product=json.loads(facts_path.read_text(encoding='utf-8-sig')) if facts_path.exists() else None
                check_product_location(path,ledger.get('productId'),product)
                if verify_projection:
                    verify_workbook(book, ledger, mode='draft')
            elif not isinstance(ledger, list):
                raise ValueError('Research ledger must be a legacy array or scoped object')
    finally:
        book.close()

def current_workbooks(root):
    # Immediate files are declared current. Explicit historical bundles live under
    # workbook-snapshots/<revision>/ and receive hash checks, never current projection.
    from workbook_paths import current_product_artifacts
    return current_product_artifacts(root,'.xlsx')


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--directory',default='.');ap.add_argument('--base');ap.add_argument('--head');args=ap.parse_args()
    root=Path(args.directory);topology=template_topology(root/'assets/listing-workbook-template.xlsx')
    from research_history import check_research_ledger_history
    check_research_ledger_history(root,args.base,args.head)
    from serial_research import check_serial_research_delivery
    serial_report=check_serial_research_delivery(root,args.base,args.head)
    check_template_inputs(root/'assets/listing-workbook-template.xlsx')
    outputs=current_workbooks(root)
    snapshot_count=check_snapshots(root)
    check_snapshot_history(root,args.base,args.head)
    scoped_checked=0
    legacy_not_assessed=0
    for path in outputs:
        manufacturer=None
        if '_LISTING_' in path.stem.upper():
            facts=json.loads((path.parent/'product-facts.json').read_text(encoding='utf-8-sig'))
            ledger_path=path.parent/'research-ledger.json'
            if ledger_path.exists() and isinstance(json.loads(ledger_path.read_text(encoding='utf-8-sig')),dict):
                from product_metadata import listing_view
                from research_export import load_ledger
                facts,_=listing_view(facts,load_ledger(ledger_path))
            identity=facts.get('identity',{})
            manufacturer=identity.get('oemBrand') if isinstance(identity,dict) else None
            if not manufacturer: raise ValueError(f'{path}: OEM brand required for warranty copy check')
        check_output(path,topology,manufacturer)
        ledger_path=path.parent/'research-ledger.json'
        if ledger_path.exists():
            if isinstance(json.loads(ledger_path.read_text(encoding='utf-8-sig')), dict):
                scoped_checked+=1
            else:
                legacy_not_assessed+=1
    from research_handoff import check_research_summaries
    summaries_checked=check_research_summaries(root)
    summary={'templateStructure':'PASS','currentOutputsStructureChecked':len(outputs),
             'serialResearch':serial_report,
             'researchSummariesFreshnessChecked':summaries_checked,
             'historicalOutputsExcluded':True,'factOrPublicationReadinessInferred':False,
             'scopedDraftProjectionChecked':scoped_checked,
             'legacyLedgerScopeNotAssessed':legacy_not_assessed,
             'historicalSnapshotBundlesHashChecked':snapshot_count}
    (root/'workbook-ci-summary.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(summary))
    if os.environ.get('GITHUB_STEP_SUMMARY'):
        with open(os.environ['GITHUB_STEP_SUMMARY'],'a',encoding='utf-8') as f:
            f.write(f'### Workbook structure scope\n\nTemplate structure PASS; {len(outputs)} current output workbooks checked. Historical workbooks excluded. No factual, warranty applicability or publication readiness inferred.\n')

if __name__=='__main__':main()
