"""Plan every template field; build and verify one coherent draft/research bundle.

Research and original copy are performed by the executing agent, not invented here.
This tool runs the existing evidence, selection, copy and saved-byte checks.
"""
import argparse
import hashlib
from html import escape
import json
from pathlib import Path
import subprocess
import sys

import openpyxl

import research_export as export
from candidate_coverage import require_current, review_projection
from research_audit import validate as audit_research

ROOT = Path(__file__).resolve().parents[1]
REPORT = 'listing-draft-bundle.json'
HTML = 'field-recommendations.html'
COPY_FIELDS = {'Item Name', 'Bullet Point', 'Product Description'}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def template_fields(template):
    book = openpyxl.load_workbook(template)
    try:
        return [{'sheet': s.title, 'attribute': s.cell(1, c).value,
                 'definition': s.cell(2, c).value}
                for s in book for c in range(2, s.max_column + 1)]
    finally:
        book.close()


def focus(attribute):
    name = attribute.casefold()
    if 'generation' in name or 'cpu' in name or 'processor' in name:
        return 'Exact installed CPU SKU and official CPU family/generation; base versus boost frequency.'
    if 'battery' in name or 'batteries' in name or 'lithium' in name:
        return 'Installed pack part/BOM and official battery table; distinguish packs, cells, Wh, chemistry and net/max weight.'
    if 'memory' in name or 'ram' in name:
        return 'Installed module/BOM, board/port variant and operating rate; module rating and MT/s are not clock MHz.'
    if 'form factor' in name:
        return 'Actual chassis mechanism and exact field Definition; normalize only to applicable observed choices.'
    return 'Read this Definition; research exact model/current sold configuration and applicable official/component documentation.'


def plan(template, data=None):
    """Diagnostic queue: retain actual Definitions and all gaps, without making facts."""
    selected = {f['id']: f for f in data['facts']} if data else {}
    mappings = {(m['sheet'], m['attribute']): m['selected_fact'] for m in data['fields']} if data else {}
    rows = []
    for field in template_fields(template):
        fact = selected.get(mappings.get((field['sheet'], field['attribute'])), {})
        evidence = fact.get('evidence', {})
        trace = evidence.get('research_trace', {})
        rows.append({**field, 'factId': fact.get('id'), 'status': fact.get('status', 'NOT_RESEARCHED'),
            'fieldClass': trace.get('field_class', 'UNCLASSIFIED_READ_DEFINITION'),
            'hasHardwareCandidate': fact.get('value_candidate') is not None,
            'hasComparativeAssessment': bool(evidence.get('candidate_assessment')),
            'recordedRoutes': trace.get('routes', []),
            'remainingRoutes': sorted({'OEM_SERIAL', 'OEM_MODEL', 'ERP', 'AMAZON', 'DEEP'} -
                                     {r['route'] for r in trace.get('routes', []) if r.get('result') != 'NOT_STARTED'}),
            'researchFocus': focus(field['attribute'])})
    return {'schemaVersion': 1, 'artifactRole': 'DRAFT_RESEARCH_WORKLIST',
        'templateSha256': digest(template), 'fieldCount': len(rows), 'fields': rows,
        'scope': 'Work queue only; no source lookup or fact confirmation is implied.'}


def readiness(data, projected):
    coverage = review_projection(data, projected, required=True)
    facts = {f['id']: f for f in data['facts']}
    fields = []
    pending = []
    for row in coverage['unresolvedFields']:
        record = projected[(row['sheet'], row['attribute'])]
        fact = facts[record['fact_id']]
        trace = fact.get('evidence', {}).get('research_trace', {})
        assessment = fact.get('evidence', {}).get('candidate_assessment')
        item = {**row, 'status': record['status'], 'source': record['source'],
                'fieldClass': trace.get('field_class'), 'researchOutcome': trace.get('outcome'),
                'hardwareAssessment': assessment}
        fields.append(item)
        if trace.get('field_class') == 'TECHNICAL' and (
                trace.get('outcome') in {'INCOMPLETE', 'BLOCKED'} or
                (not trace.get('exhausted') and fact['status'] not in {'VERIFIED', 'NOT_APPLICABLE'})):
            pending.append({'sheet': row['sheet'], 'attribute': row['attribute'],
                            'outcome': trace.get('outcome'), 'routes': trace.get('routes', [])})
    return {'coverage': coverage, 'unresolvedFields': fields, 'remainingTechnicalResearch': pending,
            'researchState': 'PARTIAL' if fields or pending else 'RECORDED_VALUES_RESOLVED',
            'sourceAuthenticationInferred': False, 'marketplaceExportAssessed': False}


def render_recommendations(data, projected):
    """All unresolved fields, including seller inputs, not just ranked candidates."""
    report = readiness(data, projected)
    rows = []
    for item in report['unresolvedFields']:
        assessment = item.get('hardwareAssessment')
        if assessment:
            from candidate_ranking import display
            options = assessment['options']
            primary = next((o for o in options if o['id'] == assessment.get('primary_id')), None)
            shown = display(primary) if primary else ' / '.join(display(o) for o in options)
        else:
            shown = str(item['candidate']) if item.get('candidate') is not None else item['decision']
        rows.append('<tr><td>' + escape(item['sheet'] + ' / ' + item['attribute']) + '</td><td><strong>'
                    + escape(shown) + '</strong><br>NOT CONFIRMED</td><td>'
                    + escape(item.get('explanation') or '') + '</td></tr>')
    return ('<!doctype html><html lang="en"><meta charset="utf-8"><title>' + escape(data['productId'])
            + ' field recommendations</title><style>body{font:16px Arial;margin:24px;color:#142639}'
            + 'table{border-collapse:collapse;width:100%}td,th{padding:12px;border:1px solid #ccd5df;'
            + 'text-align:left;vertical-align:top;overflow-wrap:anywhere}</style><h1>' + escape(data['productId'])
            + ' — Field recommendations and remaining inputs</h1><p>Research state: ' + report['researchState']
            + '. Candidates are references for review, not confirmed sellable-unit claims. Original units and limitations remain applicable.'
            + '</p><table><tr><th>Field</th><th>Preferred reference / disposition</th><th>Evidence, alternatives and missing confirmation</th></tr>'
            + ''.join(rows) + '</table></html>')


def expected_names(product_id):
    if Path(product_id).name != product_id or any(c in product_id for c in '/\\:'):
        raise ValueError('Local canonical product ID required')
    return {stage: product_id + '_' + stage + '.xlsx' for stage in ('RESEARCH_REVIEW', 'LISTING_DRAFT')}


def _context(product, template):
    product = Path(product)
    data = export.load_ledger(product/'research-ledger.json')
    require_current(data)
    if data.get('listing_draft_contract_version') != 1 or data.get('fcc_management') != 'ERP_INVENTORY':
        raise ValueError('New full draft requires listing_draft_contract_version: 1 and separate ERP FCC management')
    if digest(template) != digest(ROOT/'assets/listing-workbook-template.xlsx'):
        raise ValueError('Full draft requires the current canonical template')
    audit_research(data, required=True)
    from serial_research import validate_serial_research
    validate_serial_research(data, required=True, full_research=True)
    from research_handoff import check_product_location
    check_product_location(product/REPORT, data['productId'])
    from product_identity import product_id
    facts = read(product/'product-facts.json')
    if product_id(facts) != data['productId'] or facts.get('selectedBaseConfiguration', {}).get('id') != data['configurationId']:
        raise ValueError('Draft bundle product/configuration identity mismatch')
    book = openpyxl.load_workbook(template)
    try:
        projected = export.project(data, export.workbook_fields(book))
    finally:
        book.close()
    from amazon_selection import project_if_bound
    if 'amazon_selection' not in data:
        project_if_bound(data, projected, 'draft', required=True)
    return data, projected


def _inputs(product, template):
    revision = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    return {'workflowSha': revision, 'template': digest(template), **{n: digest(Path(product)/n) for n in
            ('research-ledger.json', 'product-facts.json', 'reference-audit.json')}}


def _write_new(path, value):
    payload = value if isinstance(value, str) else json.dumps(value, indent=2, ensure_ascii=False) + '\n'
    with Path(path).open('x', encoding='utf-8', newline='\n') as handle:
        handle.write(payload)


def build(product, template=None):
    product = Path(product)
    template = Path(template or ROOT/'assets/listing-workbook-template.xlsx')
    data, projected = _context(product, template)
    names = expected_names(data['productId'])
    targets = [n for name in names.values() for n in (name, str(Path(name).with_suffix('.research-summary.json')))] + [HTML, REPORT]
    if any((product/n).exists() for n in targets):
        raise ValueError('Bundle output already exists; preserve the prior coherent version and use a clean staged product directory')
    inputs = _inputs(product, template)
    created = []
    try:
        for stage, name in names.items():
            output = product/name
            summary = output.with_suffix('.research-summary.json')
            # Use the supported generator, including the existing trace and reference audits.
            command = [sys.executable, str(ROOT/'scripts/research_export.py'), '--ledger', str(product/'research-ledger.json'),
                '--template', str(template), '--workbook', str(output), '--write-workbook', '--stage', stage,
                '--product-id', data['productId'], '--configuration-id', data['configurationId']]
            completed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
            if completed.returncode:
                raise ValueError('Draft bundle ' + stage + ' failed: ' + completed.stderr.strip())
            for path in (output, summary):
                created.append((path, digest(path)))
        _write_new(product/HTML, render_recommendations(data, projected))
        created.append((product/HTML, digest(product/HTML)))
        if inputs != _inputs(product, template):
            raise ValueError('Research inputs changed during bundle generation')
        report = {'schemaVersion': 1, 'artifactRole': 'LISTING_DRAFT_BUNDLE', 'productId': data['productId'],
            'configurationId': data['configurationId'], 'generationState': 'DRAFT_BUNDLE_GENERATED',
            'inputs': inputs, 'workbooks': names, 'readiness': readiness(data, projected),
            'files': {p.name: sha for p, sha in created},
            'deliveryState': 'NOT_UPLOADED', 'publicationStatus': 'NOT_EXECUTED'}
        _write_new(product/REPORT, report)
        created.append((product/REPORT, digest(product/REPORT)))
        return verify(product, template)
    except BaseException:
        # Remove only these newly generated bytes; never erase old or concurrently changed files.
        for path, sha in reversed(created):
            if path.exists() and digest(path) == sha:
                path.unlink()
        raise


def verify(product, template=None):
    product = Path(product)
    template = Path(template or ROOT/'assets/listing-workbook-template.xlsx')
    data, projected = _context(product, template)
    report = read(product/REPORT)
    names = expected_names(data['productId'])
    if (report.get('schemaVersion') != 1 or report.get('artifactRole') != 'LISTING_DRAFT_BUNDLE'
            or report.get('generationState') != 'DRAFT_BUNDLE_GENERATED'
            or report.get('publicationStatus') != 'NOT_EXECUTED'
            or report.get('productId') != data['productId'] or report.get('configurationId') != data['configurationId']
            or report.get('workbooks') != names or report.get('inputs') != _inputs(product, template)
            or report.get('readiness') != readiness(data, projected)):
        raise ValueError('Draft bundle inputs, identity or readiness are stale')
    expected = {HTML, *(n for name in names.values() for n in (name, str(Path(name).with_suffix('.research-summary.json'))))}
    if set(report.get('files', {})) != expected:
        raise ValueError('Draft bundle must contain both workbooks, both summaries and all-field recommendations')
    for name in expected:
        if digest(product/name) != report['files'][name]:
            raise ValueError('Draft bundle saved bytes changed: ' + name)
    if (product/HTML).read_text(encoding='utf-8') != render_recommendations(data, projected):
        raise ValueError('Candidate recommendations differ from current ledger projection')
    for stage, name in names.items():
        result = export.run(template, product/'research-ledger.json', product/name, 'draft',
                            data['productId'], data['configurationId'], stage, check_copy=stage == 'LISTING_DRAFT')
        summary = read((product/name).with_suffix('.research-summary.json'))
        for key in ('productId', 'configurationId', 'stage', 'workbookSha256', 'ledgerSha256',
                    'templateSha256', 'productFactsSha256', 'coverageSummary'):
            if summary.get(key) != result[key]:
                raise ValueError('Draft bundle summary stale: ' + key)
        if summary.get('candidateCoverage') != review_projection(data, projected, required=True):
            raise ValueError('Draft bundle candidate summary is stale')
        from research_audit import validate_summary_trace
        validate_summary_trace(data, summary['researchTrace'])
    return {'state': 'DRAFT_BUNDLE_SAVED_BYTES_CHECK_PASS', 'productId': data['productId'],
            'files': sorted([*expected, REPORT]), 'fieldCount': len(projected),
            'readiness': report['readiness'], 'deliveryState': 'NOT_VERIFIED_BY_THIS_CHECK',
            'publicationStatus': 'NOT_EXECUTED'}


def check_delivery(product, selected, roles):
    """Guard current full-draft bundles at the existing production/delivery boundary."""
    data = read(Path(product)/'research-ledger.json')
    if data.get('listing_draft_contract_version') != 1 or not any(roles[n] == 'LISTING' for n in selected):
        return None
    result = verify(product)
    required = {*result['files'], 'research-ledger.json', 'product-facts.json', 'reference-audit.json'}
    if not required <= set(selected):
        raise ValueError('Full draft delivery omits bundle files/dependencies: ' + ', '.join(sorted(required-set(selected))))
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['plan', 'build', 'verify'])
    parser.add_argument('--product', type=Path)
    parser.add_argument('--template', type=Path, default=ROOT/'assets/listing-workbook-template.xlsx')
    parser.add_argument('--output', type=Path, help='New diagnostic worklist JSON only; never updates product facts')
    args = parser.parse_args()
    try:
        if args.output and args.command != 'plan':
            raise ValueError('--output is only for the diagnostic plan')
        if args.command == 'plan':
            ledger = args.product/'research-ledger.json' if args.product else None
            result = plan(args.template, export.load_ledger(ledger) if ledger and ledger.exists() else None)
            if args.output:
                _write_new(args.output, result)
        elif not args.product:
            raise ValueError('build/verify require --product')
        else:
            result = build(args.product, args.template) if args.command == 'build' else verify(args.product, args.template)
    except (ValueError, OSError, json.JSONDecodeError) as error:
        parser.exit(1, str(error) + '\n')
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
