"""Real workbook checks must preserve ledger-bound source-only dispositions."""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
import openpyxl
sys.path.insert(0,str(Path(__file__).parents[1]))
import research_export as export
spec = importlib.util.spec_from_file_location('candidate_workbook_checks', Path(__file__).parents[1]/'check-workbooks.py')
checks = importlib.util.module_from_spec(spec); spec.loader.exec_module(checks)


class CandidateDispositionWorkbook(unittest.TestCase):
    def test_real_checker_accepts_only_exact_scoped_disposition(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch); path=root/'review.xlsx'
            d={'schemaVersion':1,'candidate_coverage_version':1,'productId':'TEST','configurationId':'BASE','facts':[],'fields':[]}
            b=openpyxl.Workbook(); b.remove(b.active)
            for i,title in enumerate(checks.SHEETS):
                name='Memory Speed' if title=='Product Details' else 'Synthetic Field'
                fact={'id':str(i),'sheet':title,'attribute':name,'configurationId':'BASE','fact_scope':'SOLD_BASE',
                      'status':'TBD' if i==0 else 'NOT_APPLICABLE','final_value':None,'source':'Synthetic disposition source'}
                if i==0:
                    fact['evidence']={'candidate_review':{'decision':'RESEARCH_PENDING','reason':'Synthetic board branch not yet checked.',
                        'missing_confirmation':'Synthetic operating rate missing.','next_action':'Inspect synthetic target board.',
                        'owner':'Synthetic researcher','reviewed_date':'2026-10-30'}}
                d['facts'].append(fact); d['fields'].append({'sheet':title,'attribute':name,'selected_fact':str(i)})
                s=b.create_sheet(title)
                for r,label in enumerate(('Field / Attribute','Definition','Value','Status','Source','Value Candidate','Value Candidate Source'),1):s.cell(r,1,label)
                s.cell(1,2,name);s.cell(2,2,'Synthetic field definition')
            p=export.project(d,[(m['sheet'],m['attribute']) for m in d['fields']])
            for s in b:
                rec=p[(s.title,s.cell(1,2).value)]
                for r,key in ((3,'value'),(4,'status'),(5,'source'),(6,'value_candidate'),(7,'value_candidate_source')):s.cell(r,2,rec[key])
            b.save(path); b.close(); (root/'research-ledger.json').write_text(json.dumps(d))
            topology=checks.template_topology(path); checks.check_output(path,topology)
            altered=openpyxl.load_workbook(path);altered['Product Details'].cell(7,2,'Forged unsupported recommendation');altered.save(path);altered.close()
            with self.assertRaisesRegex(ValueError,'differs'):checks.check_output(path,topology)
            (root/'research-ledger.json').unlink()
            with self.assertRaisesRegex(ValueError,'scoped'):checks.check_output(path,topology)
