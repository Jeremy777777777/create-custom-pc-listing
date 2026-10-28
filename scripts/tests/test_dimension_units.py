"""Unit conversion and new-output enforcement, with synthetic data only."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import dimension_units as d
import research_export as export
import test_fcc_guide_writer as fixture
from reference_fixtures import write_audit


class DimensionUnitsTests(unittest.TestCase):
    def test_exact_conversions_ranges_and_tuple_order(self):
        for value,unit in [(254,'mm'),(25.4,'cm'),(.254,'m'),(10,'inch'),(10,'INCHES')]:
            self.assertEqual(d.to_inches(value,unit),10)
        self.assertEqual(d.to_inches(1,'ft'),12)
        self.assertEqual(d.convert_length_text('17.9–18.9 mm'),'0.7047–0.7441 inches')
        self.assertEqual(d.convert_length_text('254 x 127 x 25.4 mm'),'10 x 5 x 1 inches')
        for value,unit in [(0,'mm'),(-1,'cm'),(True,'mm'),(float('nan'),'mm'),(1,'unknown'),(1,'kg'),(.0000001,'mm')]:
            with self.subTest(value=value,unit=unit),self.assertRaises(ValueError):d.to_inches(value,unit)
        with self.assertRaises(ValueError):d.convert_length_text('254 x127')

    def test_all_numeric_pairs_and_candidates_require_inch_units(self):
        for (sheet,field),unitfield in d.PAIRS.items():
            good={(sheet,field):{'value':10},(sheet,unitfield):{'value':'Inches'}}
            self.assertEqual(d.validate_dimensions(good)['state'],'INCH_DIMENSIONS_CHECK_PASS')
            for unit in ['Millimeters','Centimeters','Feet',None]:
                bad=copy.deepcopy(good);bad[(sheet,unitfield)]['value']=unit
                with self.subTest(field=field,unit=unit),self.assertRaises(ValueError):d.validate_dimensions(bad)
            candidate={(sheet,field):{'value':None,'value_candidate':'0.7–0.74'},(sheet,unitfield):{'value':None,'value_candidate':'Inches'}}
            d.validate_dimensions(candidate)
            candidate[(sheet,unitfield)]['value_candidate']='Millimeters'
            with self.assertRaises(ValueError):d.validate_dimensions(candidate)

    def test_qualifiers_and_dimension_context_survive_conversion(self):
        pairs=[
            ('254 x 127 x 25.4 mm maximum (OLED chassis)','10 x 5 x 1 inches maximum (OLED chassis)'),
            ('254 x 127 x 25.4 mm nominal (without stand)','10 x 5 x 1 inches nominal (without stand)'),
            ('17.9–18.9 mm maximum','0.7047–0.7441 inches maximum'),
        ]
        for original,expected in pairs:
            with self.subTest(original=original):
                self.assertEqual(d.convert_length_text(original),expected)
                for slot in ('value','value_candidate'):
                    d.validate_dimensions({('Product Details','Item Dimensions L x W x Thickness'):{slot:expected}})

    def test_front_rear_maximum_profile_retains_each_position(self):
        original='254 x 127 mm; front 12.7 mm, rear 2.54 cm, maximum 1.5 inches thickness'
        expected='10 x 5 inches; front 0.5 inches, rear 1 inches, maximum 1.5 inches thickness'
        self.assertEqual(d.convert_length_text(original),expected)
        self.assertEqual(d.validate_dimension_text(expected)['representation'],'profile')
        # Front/rear are positions, not sorted range bounds; front may be taller.
        d.validate_dimension_text('10 x 5 inches; front 1 inches, rear 0.5 inches, maximum 1.5 inches thickness')
        d.validate_dimension_text('10 x 5 inches; front 0.5-0.6 inches, rear 1 inches, maximum 1.5 inches thickness')

    def test_profiles_reject_metric_residue_missing_units_and_ambiguous_labels(self):
        invalid=[
            '10 x 5 inches; front 12.7 mm, rear 1 inches, maximum 1.5 inches thickness',
            '10 x 5 inches; front 0.5, rear 1 inches, maximum 1.5 inches thickness',
            '10 x 5 inches; front 0.5 inches, front 1 inches thickness',
            '10 x 5 inches; battery 0.5 inches, rear 1 inches thickness',
            '10 x 5 inches; 0.5 inches, 1 inches thickness',
            '10 x 5 inches; front 0 inches, rear 1 inches thickness',
            '10 x 5 inches; front 1-0.5 inches, rear 1 inches thickness',
            '10 x 5 inches; front 0.5 inches, rear 1 inches, maximum 0.7 inches thickness',
            '10 x 5 inches; front 0.5 inches, rear 1 inches, minimum 0.7 inches thickness',
            '10 x 5 x 1 inches confirmed',
            '10 x 5 x 1 inches maximum (250 mm chassis)',
        ]
        for value in invalid:
            with self.subTest(value=value),self.assertRaises(ValueError):
                d.validate_dimension_text(value)
        with self.assertRaises(ValueError):
            d.convert_length_text('254 x 127 mm; front 12.7, rear 25.4 mm thickness')
        with self.assertRaisesRegex(ValueError,'inches'):
            d.validate_dimension_copy('Thickness front:12.8 mm')
        d.validate_dimension_copy('Rear 3.5mm audio jack; M.2 2280 SSD')

    def test_writer_keeps_named_inch_profile_without_mutating_the_baseline(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,_,old,ledger=fixture.FCCGuideWriterTests().fixture(root)
            fact=next(f for f in data['facts'] if f['attribute']=='Item Dimensions L x W x Thickness')
            original='254 x 127 mm; front 12.7 mm, rear 25.4 mm, maximum 38.1 mm thickness'
            fact.update(status='VERIFIED',final_value=original)
            ledger.write_text(json.dumps(data));write_audit(ledger);before=old.read_bytes();out=root/'profile.xlsx'
            with self.assertRaises(ValueError):
                export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            self.assertFalse(out.exists());self.assertEqual(old.read_bytes(),before)
            fact['final_value']=d.convert_length_text(original)
            fact['evidence']={'original_geometry':original,'scope':'Synthetic position-labelled design, not a unit measurement'}
            ledger.write_text(json.dumps(data));write_audit(ledger)
            export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            import openpyxl
            book=openpyxl.load_workbook(out);sheet=book['Product Details']
            col=next(c.column for c in sheet[1] if c.value==fact['attribute'])
            self.assertEqual(sheet.cell(3,col).value,'10 x 5 inches; front 0.5 inches, rear 1 inches, maximum 1.5 inches thickness')
            self.assertEqual(sheet.cell(4,col).value,'VERIFIED');book.close()
            self.assertEqual(old.read_bytes(),before)

    def test_composite_and_bezel_strings_need_explicit_inches(self):
        for key in d.TEXT_FIELDS:
            d.validate_dimensions({key:{'value':'10 x 5 x 1 inches'}})
            for value in ['254 x127 x25.4 mm','10 x5 x1',2,True,'0 inches']:
                with self.subTest(key=key,value=value),self.assertRaises(ValueError):d.validate_dimensions({key:{'value':value}})
        d.validate_dimensions({})  # Unknown dimensions never become an invented number.

    def test_customer_size_copy_and_technical_designations(self):
        for text in ['Dimensions:254 x127 x25.4 mm','Height:18 mm','Width is 25 cm','Depth 1 foot']:
            with self.subTest(text=text),self.assertRaises(ValueError):d.validate_dimension_copy(text)
        for text in ['Dimensions:10 x5 x1 inches','M.2 2280 SSD;3.5mm audio;1920 x1200 pixels','Weight:1.59kg;60Hz']:
            d.validate_dimension_copy(text)

    def test_descriptive_geometry_is_checked_in_formal_and_candidate_rows(self):
        for field,text in [('Notebook Pointing Device Description','Touchpad,115 x80 mm'),('Touchpad Feature','115 mm x 80 mm touchpad'),('Keyboard Description','Key travel:1.5 mm'),('Keyboard Description','1.5mm key travel')]:
            for slot in ('value','value_candidate'):
                with self.subTest(field=field,slot=slot),self.assertRaisesRegex(ValueError,'inches'):
                    d.validate_dimensions({('Product Details',field):{slot:text}})
        data={('Product Details','Touchpad Feature'):{'value':'Touchpad,4.5276 x3.1496 inches','source':'Original115 x80mm','value_candidate_source':'Original115 mm x80 mm'}}
        before=copy.deepcopy(data);d.validate_dimensions(data);self.assertEqual(data,before)

    def test_named_connector_count_does_not_hide_explicit_geometry(self):
        for text in ['1 x3.5mm audio jack','2 x 3.5 mm headphone jacks','1x3.5mm (headphone/microphone) combo;M.2 2280;2x2 Wi-Fi','1 x 3.5mm global headset jack, input and output','1 x3.5mm universal audio jack']:
            d.validate_dimension_copy(text)
            d.validate_dimensions({('Product Details','Hardware Interface'):{'value':text}})
        for text in ['Dimensions:2 x3.5mm audio jack body','Width:2 x3.5mm jack body','Touchpad115 mm by80 mm','Height:0.25 m']:
            with self.subTest(text=text),self.assertRaisesRegex(ValueError,'inches'):d.validate_dimension_copy(text)

    def test_writer_rejects_metric_description_without_mutating_original(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,_,old,ledger=fixture.FCCGuideWriterTests().fixture(root)
            fact=next(f for f in data['facts'] if f['attribute']=='Touchpad Feature')
            fact.update(status='VERIFIED',final_value='Touchpad,115 x80 mm')
            ledger.write_text(json.dumps(data));write_audit(ledger);before=old.read_bytes();out=root/'blocked.xlsx'
            with self.assertRaisesRegex(ValueError,'inches'):export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            self.assertFalse(out.exists());self.assertEqual(old.read_bytes(),before)
            fact['evidence']={'original_geometry':'115 x80 mm','conversion':'Each axis divided by25.4'}
            fact['final_value']='Touchpad,'+d.convert_length_text('115 x80 mm')
            ledger.write_text(json.dumps(data));write_audit(ledger);export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            import openpyxl
            book=openpyxl.load_workbook(out);sheet=book['Product Details'];col=next(c.column for c in sheet[1] if c.value=='Touchpad Feature')
            self.assertEqual(sheet.cell(3,col).value,'Touchpad,4.5276 x3.1496 inches');book.close()

    def test_real_writer_rejects_metric_and_accepts_converted_ledger(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,_,old,ledger=fixture.FCCGuideWriterTests().fixture(root)
            length=next(f for f in data['facts'] if f['attribute']=='Item Length Horizontal Edge')
            unit=next(f for f in data['facts'] if f['attribute']=='Item Length Unit')
            length.update(status='VERIFIED',final_value=254)
            unit.update(status='VERIFIED',final_value='Millimeters')
            ledger.write_text(json.dumps(data));write_audit(ledger)
            before=old.read_bytes()
            with self.assertRaisesRegex(ValueError,'Inches'):export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,root/'blocked.xlsx')
            self.assertFalse((root/'blocked.xlsx').exists());self.assertEqual(old.read_bytes(),before)
            length['evidence']={'original_value':254,'original_unit':'mm','conversion':'254 /25.4'}
            length['final_value']=d.to_inches(254,'mm');unit['final_value']='Inches'
            ledger.write_text(json.dumps(data));write_audit(ledger)
            out=root/'new.xlsx';export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            import openpyxl
            book=openpyxl.load_workbook(out);sheet=book['Product Details']
            col=next(c.column for c in sheet[1] if c.value=='Item Length Horizontal Edge')
            self.assertEqual(sheet.cell(3,col).value,10);book.close()


if __name__=='__main__':unittest.main()
