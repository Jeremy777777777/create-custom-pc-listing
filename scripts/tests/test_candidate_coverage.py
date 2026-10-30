"""Prevent silent candidate omissions while preserving honest unknowns and history."""
import copy
from io import BytesIO
from datetime import date
from pathlib import Path
import sys
import unittest
import openpyxl
import json
import subprocess
import tempfile

sys.path.insert(0, str(Path(__file__).parents[1]))
import candidate_coverage as coverage
import research_export as export


class CandidateCoverageTests(unittest.TestCase):
    key = ('Safety&Compliance', 'Battery Weight')

    def ledger(self):
        return {'schemaVersion': 1, 'candidate_coverage_version': 1,
                'productId': 'TEST-PC', 'configurationId': 'BASE',
                'facts': [{'id': 'weight', 'sheet': self.key[0], 'attribute': self.key[1],
                           'configurationId': 'BASE', 'fact_scope': 'SOLD_BASE',
                           'status': 'TBD', 'final_value': None, 'source': 'Exact net weight not measured',
                           'evidence': {}}],
                'fields': [{'sheet': self.key[0], 'attribute': self.key[1], 'selected_fact': 'weight'}]}

    def pending(self):
        return {'decision': 'RESEARCH_PENDING', 'reason': 'Exact installed battery part weight is not yet checked.',
                'missing_confirmation': 'Net weight of installed part; OEM maximum differs.',
                'next_action': 'Inspect the matched official battery part record.',
                'owner': 'Product researcher', 'reviewed_date': date.today().isoformat()}

    def test_silent_gap_fails_but_history_remains_readable(self):
        d = self.ledger()
        with self.assertRaisesRegex(ValueError, 'candidate_review'):
            export.project(d, [self.key])
        del d['candidate_coverage_version']
        self.assertIsNone(export.project(d, [self.key])[self.key]['value_candidate_source'])
        with self.assertRaisesRegex(ValueError, 'candidate_coverage_version'):
            coverage.require_current(d)

    def test_new_generation_cli_blocks_legacy_gap_before_writing(self):
        d = self.ledger(); del d['candidate_coverage_version']
        d['fcc_management'] = 'ERP_INVENTORY'
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch); ledger = root/'research-ledger.json'; out = root/'draft.xlsx'
            ledger.write_text(json.dumps(d), encoding='utf-8')
            result = subprocess.run([sys.executable, str(export.ROOT/'scripts/research_export.py'),
                '--template', str(export.ROOT/'assets/listing-workbook-template.xlsx'),
                '--ledger', str(ledger), '--workbook', str(out), '--write-workbook'],
                capture_output=True, text=True, encoding='utf-8')
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('candidate_coverage_version', result.stderr)
            self.assertFalse(out.exists())

    def test_maximum_is_visible_candidate_never_exact_value(self):
        d = self.ledger(); f = d['facts'][0]
        f.update(value_candidate=190, value_candidate_source='https://example.test/oem-battery: maximum 0.19 kg = 190 grams')
        f['evidence'].update(candidate_basis='OEM_MAXIMUM',
                             candidate_rationale='OEM maximum for matching 3-cell pack, not measured net weight.',
                             candidate_missing_confirmation='Exact installed pack net weight remains unverified.')
        p = export.project(d, [self.key])
        self.assertIsNone(p[self.key]['value'])
        self.assertEqual(p[self.key]['value_candidate'], 190)
        self.assertIn('maximum', p[self.key]['value_candidate_source'])
        self.assertIn('QUALIFIED_REFERENCE', p[self.key]['value_candidate_source'])
        self.assertNotIn('MOST_LIKELY', p[self.key]['value_candidate_source'])
        self.assertEqual(coverage.review_projection(d, p)['counts']['SELECTED'], 1)
        with self.assertRaisesRegex(ValueError, 'unresolved'):
            export.project(d, [self.key], 'export')

    def test_empty_candidate_has_visible_actionable_reason(self):
        d = self.ledger(); d['facts'][0]['evidence']['candidate_review'] = self.pending()
        p = export.project(d, [self.key]); r = p[self.key]
        self.assertIsNone(r['value_candidate'])
        for expected in ('RESEARCH_PENDING', 'Net weight', 'Inspect', 'Product researcher'):
            self.assertIn(expected, r['value_candidate_source'])
        self.assertEqual(coverage.review_projection(d, p)['state'], 'PARTIAL')
        for name in ('reason', 'missing_confirmation', 'next_action', 'owner', 'reviewed_date'):
            altered = copy.deepcopy(d)
            altered['facts'][0]['evidence']['candidate_review'][name] = 'TBD'
            with self.subTest(name=name), self.assertRaises(ValueError):
                export.project(altered, [self.key])

    def test_no_support_cannot_hide_unfinished_research(self):
        d = self.ledger(); review = self.pending(); review['decision'] = 'NO_SUPPORT'
        d['facts'][0]['evidence']['candidate_review'] = review
        with self.assertRaisesRegex(ValueError, 'exhausted'):
            export.project(d, [self.key])

    def test_tie_needs_sources_and_shows_alternatives(self):
        d = self.ledger(); review = self.pending(); review['decision'] = 'TIED'
        d['facts'][0]['evidence']['candidate_review'] = review
        with self.assertRaisesRegex(ValueError, 'sourced alternatives'):
            export.project(d, [self.key])
        review['alternatives'] = [{'value': 180, 'source': 'https://example.test/part-a'},
                                  {'value': 190, 'source': 'https://example.test/part-b'}]
        self.assertIn('180 [https://example.test/part-a]', export.project(d, [self.key])[self.key]['value_candidate_source'])

    def test_saved_gap_explanation_cannot_be_dropped(self):
        d = self.ledger(); d['facts'][0]['evidence']['candidate_review'] = self.pending()
        p = export.project(d, [self.key])[self.key]
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = self.key[0]
        for row, label in enumerate(('Field / Attribute', 'Definition', 'Value', 'Status', 'Source',
                                     'Value Candidate', 'Value Candidate Source'), 1):
            sheet.cell(row, 1, label)
        sheet.cell(1, 2, self.key[1]); sheet.cell(2, 2, 'Net weight')
        for row, name in ((3, 'value'), (4, 'status'), (5, 'source'), (6, 'value_candidate'), (7, 'value_candidate_source')):
            sheet.cell(row, 2, p[name])
        saved = BytesIO(); book.save(saved); book.close(); saved.seek(0)
        book = openpyxl.load_workbook(saved); sheet = book.active
        export.verify_workbook(book, d)
        sheet.cell(7, 2).value = None
        with self.assertRaisesRegex(ValueError, 'differs'):
            export.verify_workbook(book, d)
        book.close()


if __name__ == '__main__':
    unittest.main()
