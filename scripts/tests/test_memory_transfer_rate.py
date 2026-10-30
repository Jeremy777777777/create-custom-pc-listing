"""Synthetic research-rate versus observed Amazon-clock boundary regressions."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
sys.path.insert(0,str(Path(__file__).parents[1]))
sys.path.insert(0,str(Path(__file__).parent))
import research_export as export
from listing_content import validate_listing
import test_research_export as fixtures
from reference_fixtures import write_audit


class MemoryTransferRateTests(unittest.TestCase):
    def fixture(self,root):
        book,data,_,ledger=fixtures.ResearchExportTests().full_fixture(root);book.close()
        for field,value in (('Memory Speed',8533),('Memory Speed Unit','MT/s')):
            fact=next(f for f in data['facts'] if f['sheet']=='Product Details' and f['attribute']==field)
            fact.update(status='VERIFIED',final_value=value)
        review=next(f for f in data['amazon_selection']['fields'] if f['attribute']=='Memory Speed Unit')
        review.update(decision='APPROXIMATE',referenceOptionCode='MHz',reason='Synthetic only observed unit comparison',
                      differences='MT/s transfer rate is not MHz clock; no conversion',missingConfirmation='Correct backend unit or measured clock required')
        data['amazon_selection']['catalogReference']='https://example.test/synthetic-options'
        ledger.write_text(json.dumps(data),encoding='utf-8');write_audit(ledger)
        product=json.loads((root/'product-facts.json').read_text())
        factual=copy.deepcopy(data);factual.pop('amazon_selection')
        projected=export.project(factual,[(m['sheet'],m['attribute']) for m in factual['fields']])
        return data,ledger,product,projected

    def test_saved_bound_draft_retains_transfer_rate_and_withholds_numeric(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);data,ledger,product,_=self.fixture(root);output=root/'TEST_LISTING_DRAFT.xlsx'
            from test_customization_policy import configure
            configure(data,product,False,'LPDDR5X')
            ledger.write_text(json.dumps(data),encoding='utf-8');write_audit(ledger);before=ledger.read_bytes()
            export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,output,stage='LISTING_DRAFT')
            result=export.run(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,output,'draft',stage='LISTING_DRAFT',check_copy=True)
            self.assertEqual(result['result'],'DRAFT_PROJECTION_PASS');self.assertEqual(ledger.read_bytes(),before)
            book=openpyxl.load_workbook(output);sheet=book['Product Details']
            col={c.value:c.column for c in sheet[1]}
            self.assertIsNone(sheet.cell(3,col['Memory Speed']).value)
            self.assertIsNone(sheet.cell(6,col['Memory Speed']).value)
            self.assertIn('8533',sheet.cell(5,col['Memory Speed']).value)
            self.assertIsNone(sheet.cell(3,col['Memory Speed Unit']).value)
            self.assertEqual(sheet.cell(6,col['Memory Speed Unit']).value,'MHz')
            sheet.cell(3,col['Memory Speed']).value=8533
            with self.assertRaisesRegex(ValueError,'differs'):export.verify_workbook(book,data)
            book.close()
            with self.assertRaises(ValueError):export.project(data,[(m['sheet'],m['attribute']) for m in data['fields']],'export')

    def test_writer_blocks_inferred_clock_before_creating_a_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);data,ledger,_,_=self.fixture(root)
            for name,value in (('Memory Clock Speed',4.2665),('Memory Clock Speed Unit','GHz')):
                f=next(f for f in data['facts'] if f['sheet']=='Product Details' and f['attribute']==name)
                f.update(status='INPUT_UNVERIFIED',final_value=None,value_candidate=value,
                         value_candidate_source='https://example.test/synthetic-memory')
                f['evidence']={'candidate_rationale':'DDR8533MT/s divided by2 suggests4.2665GHz clock',
                               'candidate_missing_confirmation':'Actual runtime clock unmeasured'}
            ledger.write_text(json.dumps(data));write_audit(ledger)
            output=root/'BLOCKED.xlsx';before=ledger.read_bytes()
            with self.assertRaisesRegex(ValueError,'transfer-to-clock'):
                export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,output)
            self.assertFalse(output.exists());self.assertEqual(ledger.read_bytes(),before)

    def test_known_old_candidate_can_be_read_as_history_but_never_projected(self):
        from research_audit import legacy_reuse
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);data,_,_,_=self.fixture(root)
            changed=[]
            for name,value in (('Memory Clock Speed',4.2665),('Memory Clock Speed Unit','GHz')):
                f=next(f for f in data['facts'] if f['sheet']=='Product Details' and f['attribute']==name)
                f.update(status='INPUT_UNVERIFIED',final_value=None,value_candidate=value,
                         value_candidate_source='https://example.test/synthetic-memory')
                f['evidence']={'candidate_rationale':'DDR8533MT/s divided by2 suggests4.2665GHz clock',
                               'candidate_missing_confirmation':'Actual runtime unmeasured'}
                changed.append(f['id'])
            raw=json.dumps(data).encode();current=copy.deepcopy(data)
            for f in current['facts']:
                if f['id'] in changed:
                    f['value_candidate']=8533 if f['attribute']=='Memory Clock Speed' else 'MT/s'
            retained=next(f['id'] for f in data['facts'] if f['attribute']=='Model Number')
            review=dict(productId=data['productId'],configurationId=data['configurationId'],
                previousLedgerSha256=hashlib.sha256(raw).hexdigest(),baselineReference='https://example.test/archive',
                reviewer='Synthetic reviewer',reviewed_date='2026-01-01',impact_review='Fix clock candidates only',
                fields=[dict(factId=retained,basis='Whole unchanged identity record',dependencyFactIds=[])])
            legacy_reuse(current,[retained],raw,review)
            self.assertEqual(json.dumps(data).encode(),raw)
            with self.assertRaisesRegex(ValueError,'transfer-to-clock'):
                export.project(data,[(m['sheet'],m['attribute']) for m in data['fields']])
            changed_history=copy.deepcopy(current)
            next(f for f in changed_history['facts'] if f['id']==retained)['source']='Changed source'
            with self.assertRaisesRegex(ValueError,'changed'):
                legacy_reuse(changed_history,[retained],raw,review)

    def test_research_exception_is_opt_in_and_does_not_hide_wrong_cpu(self):
        with tempfile.TemporaryDirectory() as tmp:
            _,_,product,projected=self.fixture(Path(tmp));before=copy.deepcopy(projected)
            with self.assertRaisesRegex(ValueError,'unsupported unit'):validate_listing(projected,product)
            validate_listing(projected,product,allow_research_transfer_rate=True);self.assertEqual(projected,before)
            projected[('Product Details','CPU Model Number')]['value']='1255U'
            with self.assertRaisesRegex(ValueError,'CPU'):validate_listing(projected,product,allow_research_transfer_rate=True)

    def test_transfer_rate_is_not_allowed_for_clock_fields_or_invalid_numbers(self):
        with tempfile.TemporaryDirectory() as tmp:
            _,_,product,projected=self.fixture(Path(tmp))
            for value in (-1,float('inf'),float('nan')):
                changed=copy.deepcopy(projected);changed[('Product Details','Memory Speed')]['value']=value
                with self.subTest(value=value),self.assertRaisesRegex(ValueError,'positive numeric'):validate_listing(changed,product,allow_research_transfer_rate=True)
            projected[('Product Details','Memory Speed')]['value']=None;projected[('Product Details','Memory Speed Unit')]['value']=None
            projected[('Product Details','Memory Clock Speed')]={'value':8533,'status':'VERIFIED'}
            projected[('Product Details','Memory Clock Speed Unit')]={'value':'MT/s','status':'VERIFIED'}
            with self.assertRaisesRegex(ValueError,'unsupported unit'):validate_listing(projected,product,allow_research_transfer_rate=True)

    def test_research_rate_cannot_relabel_customer_copy_as_clock(self):
        with tempfile.TemporaryDirectory() as tmp:
            _,_,product,projected=self.fixture(Path(tmp))
            for claim in ('LPDDR5X at8533MHz','4266.5MHz RAM','Memory speed: 8400 MT/s'):
                changed=copy.deepcopy(projected);changed[('Product Details','Product Description')]['value']+=' '+claim
                with self.subTest(claim=claim),self.assertRaisesRegex(ValueError,'memory transfer rate'):validate_listing(changed,product,allow_research_transfer_rate=True)
            projected[('Product Details','Product Description')]['value']+=' LPDDR5X at8533MT/s'
            validate_listing(projected,product,allow_research_transfer_rate=True)


if __name__=='__main__':unittest.main()
