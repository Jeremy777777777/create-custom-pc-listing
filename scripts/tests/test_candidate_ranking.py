"""Ranking must be evidenced, unit-preserving, and visible in delivered review bytes."""
from io import BytesIO
from pathlib import Path
import copy
import sys
import unittest
import openpyxl
sys.path.insert(0, str(Path(__file__).parents[1]))
import candidate_ranking as ranking
import candidate_coverage as coverage


class RankedCandidates(unittest.TestCase):
    def fixture(self):
        review = {'version': 1, 'decision': 'RECOMMENDED', 'primary_id': 'fast',
                  'comparison': 'Synthetic same-variant lead favors fast; actual board is unverified.',
                  'missing_confirmation': 'Exact board and operating transfer rate.', 'reviewed_date': '2026-10-30',
                  'options': [{'id': 'fast', 'value': 5200, 'unit': 'MT/s', 'source': 'https://example.test/fast',
                               'support': 'Synthetic documented branch and same-variant lead',
                               'limitations': 'Lead is lower confidence; not an installed-part measurement', 'confidence': 'LOW'},
                              {'id': 'slow', 'value': 4400, 'unit': 'MT/s', 'source': 'https://example.test/slow',
                               'support': 'Synthetic alternative board branch',
                               'limitations': 'Exact board unknown', 'confidence': 'LOW'}]}
        f = {'id': 'speed', 'attribute': 'Memory Speed', 'status': 'INPUT_UNVERIFIED',
             'value_candidate': 5200, 'evidence': {'candidate_assessment': review}}
        d = {'productId': 'TEST', 'candidate_coverage_version': 2, 'facts': [f]}
        p = {('Product Details', 'Memory Speed'): {'fact_id': 'speed', 'value': None,
             'status': 'INPUT_UNVERIFIED', 'value_candidate': None, 'value_candidate_source': None}}
        return d, p

    def test_unit_mismatch_keeps_hardware_recommendation_visible_without_relabelling(self):
        d, p = self.fixture(); ranking.apply(d, p)
        field = next(iter(p.values()))
        self.assertIsNone(field['value']); self.assertIsNone(field['value_candidate'])
        self.assertIn('5200 MT/s (LOW confidence; NOT CONFIRMED)', field['value_candidate_source'])
        self.assertIn('4400 MT/s', field['value_candidate_source'])
        self.assertNotIn('5200 MHz', field['value_candidate_source'])
        self.assertEqual(coverage.review_projection(d, p)['counts']['HARDWARE_REFERENCE_WITHHELD_FROM_AMAZON_NUMBER'], 1)
        html = ranking.render_html(d, p)
        self.assertIn('<strong>5200 MT/s</strong>', html)
        self.assertIn('NOT CONFIRMED', html)

    def test_primary_mismatch_missing_comparison_and_invented_confidence_fail(self):
        d, _ = self.fixture()
        for mutate in (lambda r: r.update(primary_id='absent'), lambda r: r.update(comparison=''),
                       lambda r: r['options'][0].update(confidence='90%'),
                       lambda r: r['options'][0].update(value=5600),
                       lambda r: r['options'][0].update(source='')):
            f = copy.deepcopy(d['facts'][0]); mutate(f['evidence']['candidate_assessment'])
            with self.assertRaises(ValueError): ranking.assess(f)

    def test_genuine_tie_has_no_forced_primary(self):
        d, p = self.fixture(); r = d['facts'][0]['evidence']['candidate_assessment']
        r.update(decision='TIED', primary_id=None)
        ranking.apply(d, p)
        self.assertIn('TIED HARDWARE CANDIDATES: 5200 MT/s / 4400 MT/s', next(iter(p.values()))['value_candidate_source'])
        r['primary_id'] = 'fast'
        with self.assertRaises(ValueError): ranking.assess(d['facts'][0])

    def test_new_contract_rejects_old_marker_and_uncompared_candidate(self):
        with self.assertRaises(ValueError): coverage.require_current({'candidate_coverage_version': 1})
        coverage.require_current({'candidate_coverage_version': 2})
        with self.assertRaisesRegex(ValueError, 'comparative'):
            ranking.assess({'attribute': 'Memory Speed', 'value_candidate': 5200}, required=True)

    def test_excel_saved_bytes_contain_qualified_recommendation(self):
        d, p = self.fixture(); ranking.apply(d, p)
        b = openpyxl.Workbook(); s = b.active
        s.cell(7, 2, next(iter(p.values()))['value_candidate_source'])
        output = BytesIO(); b.save(output); b.close(); output.seek(0)
        saved = openpyxl.load_workbook(output)
        self.assertIn('5200 MT/s', saved.active.cell(7, 2).value)
        self.assertIsNone(saved.active.cell(6, 2).value)
        saved.close()
