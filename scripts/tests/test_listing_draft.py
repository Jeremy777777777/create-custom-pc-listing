"""Exercise the real two-workbook pipeline with explicitly synthetic research."""
import copy
import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).parents[1]))
import listing_draft as draft
import research_export as export
import test_research_export as fixtures
from test_customization_policy import configure
from test_release_checks import trace_for
import test_candidate_ranking as ranked_fixtures
from reference_fixtures import write_audit
import fcc_inventory


class DraftBundleTests(unittest.TestCase):
    def fixture(self, root):
        book, data, previous, old_ledger = fixtures.ResearchExportTests().full_fixture(root)
        book.close(); previous.unlink(); old_ledger.unlink()
        configure(data, json.loads((root/'product-facts.json').read_text()))
        data.update(candidate_coverage_version=2, listing_draft_contract_version=1, fcc_management='ERP_INVENTORY')
        external = next(f for f in data['facts'] if f['attribute'] == 'FCC ID')
        for key in ('evidence', 'evidence_scope', 'authorization_route'):
            external.pop(key, None)
        external.update(status='TBD', final_value=None, source=fcc_inventory.DISPOSITION)
        ranked, _ = ranked_fixtures.RankedCandidates().fixture()
        speed = next(f for f in data['facts'] if f['attribute'] == 'Memory Speed')
        speed.update(status='INPUT_UNVERIFIED', final_value=None, value_candidate=5200,
                     value_candidate_source='https://example.test/fast ; https://example.test/slow',
                     evidence=copy.deepcopy(ranked['facts'][0]['evidence']))
        speed['evidence'].update(candidate_rationale='Synthetic same-variant lead; actual board unverified',
                                candidate_missing_confirmation='Synthetic board and actual transfer rate')
        unit = next(f for f in data['facts'] if f['attribute'] == 'Memory Speed Unit')
        unit.update(status='INPUT_UNVERIFIED', final_value=None, value_candidate='MT/s',
                    value_candidate_source='https://example.test/fast', evidence={
                        'candidate_rationale':'Synthetic official transfer-rate quantity',
                        'candidate_missing_confirmation':'Synthetic exact operating rate',
                        'candidate_assessment':{'version':1,'decision':'RECOMMENDED','primary_id':'unit',
                            'comparison':'Synthetic source expressly uses transfer rate; clock is not inferred.',
                            'missing_confirmation':'Actual operating rate', 'reviewed_date':'2026-10-30',
                            'options':[{'id':'unit','value':'MT/s','source':'https://example.test/fast',
                                'support':'Synthetic transfer-rate label','limitations':'Not a clock frequency','confidence':'HIGH'}]}})
        from amazon_selection import read_catalog, context_for
        catalog, _ = read_catalog()
        context = context_for(catalog, data['amazon_selection']['context'])
        control = next(c for c in context['controls'] if c.get('attribute') == 'Memory Speed Unit')
        option = next(o for o in control['options'] if o['label'] == 'MHz')
        binding = next(f for f in data['amazon_selection']['fields'] if f['attribute'] == 'Memory Speed Unit')
        binding.update(decision='APPROXIMATE', referenceOptionCode=option['code'],
            reason='Synthetic closest observed rate field unit', differences='MT/s and MHz are different quantities',
            missingConfirmation='Compatible original transfer-rate unit unavailable in synthetic control')
        data['amazon_selection']['catalogReference']='https://example.test/synthetic-ui-options'
        seller = next(f for f in data['facts'] if f['sheet'] == 'Offer')
        seller.update(status='TBD', final_value=None, evidence={'candidate_review':{
            'decision':'SELLER_INPUT_REQUIRED','reason':'Synthetic seller has deferred this account field.',
            'missing_confirmation':'Synthetic account input', 'next_action':'Obtain seller account input',
            'owner':'Synthetic seller', 'reviewed_date':'2026-10-30'}})
        for f in data['facts']:
            if f is external:
                continue
            f.setdefault('evidence', {})['research_trace'] = trace_for(f)
        seller['evidence']['research_trace']={'version':1,'field_class':'SELLER_INPUT',
            'reason':'Synthetic account input is not a public hardware fact','basis':'Synthetic seller deferral'}
        ledger = root/'research-ledger.json'
        ledger.write_text(json.dumps(data), encoding='utf-8'); write_audit(ledger)
        return data, ledger

    def test_full_pipeline_writes_all_fields_candidates_and_both_stages(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch); data, ledger=self.fixture(root)
            original=ledger.read_bytes()
            result=draft.build(root)
            self.assertEqual(result['state'],'DRAFT_BUNDLE_SAVED_BYTES_CHECK_PASS')
            self.assertEqual(result['fieldCount'],len(draft.template_fields(export.ROOT/'assets/listing-workbook-template.xlsx')))
            self.assertEqual(result['readiness']['researchState'],'PARTIAL')
            self.assertTrue(result['readiness']['remainingTechnicalResearch'])
            self.assertEqual(ledger.read_bytes(),original)
            for name in draft.expected_names('TEST-PC').values():
                b=openpyxl.load_workbook(root/name)
                speed=next(c.column for c in b['Product Details'][1] if c.value=='Memory Speed')
                s=b['Product Details']
                self.assertIsNone(s.cell(3,speed).value)
                self.assertIsNone(s.cell(6,speed).value)
                self.assertIn('5200 MT/s',s.cell(7,speed).value)
                self.assertIn('4400 MT/s',s.cell(7,speed).value)
                self.assertNotIn('2600 MHz',s.cell(7,speed).value)
                b.close()
            html=(root/draft.HTML).read_text(encoding='utf-8')
            self.assertIn('<strong>5200 MT/s</strong>',html)
            self.assertIn('SELLER_INPUT_REQUIRED',html)
            self.assertEqual(len(draft.verify(root)['files']),6)
            roles={name:'LISTING' if 'LISTING_DRAFT' in name else 'QA' for name in result['files']}
            roles.update({'research-ledger.json':'RESEARCH','product-facts.json':'RESEARCH','reference-audit.json':'QA'})
            self.assertIsNotNone(draft.check_delivery(root,list(roles),roles))
            with self.assertRaisesRegex(ValueError,'omits'):
                draft.check_delivery(root,[n for n in roles if n!=draft.HTML],roles)

    def test_missing_comparison_or_silent_blank_cannot_generate(self):
        for mutation in ('comparison','blank'):
            with self.subTest(mutation=mutation),tempfile.TemporaryDirectory() as scratch:
                root=Path(scratch);data,ledger=self.fixture(root)
                speed=next(f for f in data['facts'] if f['attribute']=='Memory Speed')
                if mutation=='comparison':speed['evidence'].pop('candidate_assessment')
                else:
                    speed.pop('value_candidate');speed.pop('value_candidate_source')
                    speed['evidence'].pop('candidate_assessment')
                    speed['evidence']['research_trace']['outcome']='INCOMPLETE'
                ledger.write_text(json.dumps(data));write_audit(ledger)
                with self.assertRaises(ValueError):draft.build(root)
                self.assertFalse((root/draft.REPORT).exists())
                self.assertFalse(list(root.glob('*.xlsx')))

    def test_missing_copy_fails_and_removes_only_new_bundle_outputs(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,ledger=self.fixture(root)
            title=next(f for f in data['facts'] if f['attribute']=='Model Name')
            title.update(status='NOT_APPLICABLE',final_value=None)
            title['evidence']['research_trace']=trace_for(title)
            ledger.write_text(json.dumps(data));write_audit(ledger)
            (root/'unrelated.txt').write_text('retain')
            with self.assertRaisesRegex(ValueError,'Model Name'):draft.build(root)
            self.assertEqual((root/'unrelated.txt').read_text(),'retain')
            self.assertTrue(ledger.exists());self.assertFalse(list(root.glob('*.xlsx')))

    def test_existing_bundle_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);self.fixture(root);draft.build(root)
            before={p.name:draft.digest(p) for p in root.iterdir() if p.is_file()}
            with self.assertRaisesRegex(ValueError,'already exists'):draft.build(root)
            self.assertEqual(before,{p.name:draft.digest(p) for p in root.iterdir() if p.is_file()})

    def test_single_listing_writer_checks_core_copy_without_optional_flag(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,ledger=self.fixture(root)
            model=next(f for f in data['facts'] if f['attribute']=='Model Name')
            model.update(status='NOT_APPLICABLE',final_value=None)
            model['evidence']['research_trace']=trace_for(model)
            ledger.write_text(json.dumps(data));write_audit(ledger)
            with self.assertRaisesRegex(ValueError,'Model Name'):
                export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,
                                      root/'single.xlsx',stage='LISTING_DRAFT')
            self.assertFalse((root/'single.xlsx').exists())

    def test_new_cli_cannot_skip_the_full_draft_contract(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);data,ledger=self.fixture(root)
            data.pop('listing_draft_contract_version')
            ledger.write_text(json.dumps(data));write_audit(ledger)
            error=io.StringIO()
            with patch.object(sys,'argv',['research_export','--ledger',str(ledger),'--workbook',str(root/'single.xlsx'),
                '--write-workbook','--stage','LISTING_DRAFT']),contextlib.redirect_stderr(error),self.assertRaises(SystemExit):
                export.main()
            self.assertIn('listing_draft_contract_version',error.getvalue())
            self.assertFalse((root/'single.xlsx').exists())

    def test_tampered_candidate_and_missing_stage_fail_readback(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);self.fixture(root);draft.build(root)
            path=root/'TEST-PC_LISTING_DRAFT.xlsx'
            b=openpyxl.load_workbook(path)
            col=next(c.column for c in b['Product Details'][1] if c.value=='Memory Speed')
            b['Product Details'].cell(7,col,'');b.save(path);b.close()
            with self.assertRaisesRegex(ValueError,'saved bytes'):draft.verify(root)
            (root/'TEST-PC_RESEARCH_REVIEW.xlsx').unlink()
            with self.assertRaises((ValueError, OSError)):draft.verify(root)

    def test_template_plan_enumerates_definitions_without_inventing_facts(self):
        report=draft.plan(export.ROOT/'assets/listing-workbook-template.xlsx')
        self.assertEqual(report['fieldCount'],294)
        self.assertTrue(all(r['definition'] and r['status']=='NOT_RESEARCHED' for r in report['fields']))
        self.assertEqual(len({(r['sheet'],r['attribute']) for r in report['fields']}),294)
        self.assertFalse(any(r['hasHardwareCandidate'] for r in report['fields']))


if __name__ == '__main__':unittest.main()
