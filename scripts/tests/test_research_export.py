"""Guard sold-vs-factory scope and draft-vs-export boundaries using synthetic facts."""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import research_export as r
from warranty_policy import render


def fcc_evidence(identifier='TEST-ID'):
    return {'lookup_outcome': 'UNIT_CONFIRMED', 'module_model': 'SYNTHETIC-MODULE',
            'installation_basis': 'Synthetic unit label', 'replacement_check': 'Synthetic no changes record',
            'fcc_official_url': 'https://apps.fcc.gov/oetcf/eas/reports/GenericSearch.cfm',
            'reviewed_date': '2026-10-16', 'fcc_record_id': identifier, 'fcc_record_model': 'SYNTHETIC-MODULE',
            'fcc_applicant': 'Synthetic applicant', 'fcc_grant_date': '2026-01-01', 'fcc_access_result': 'REVIEWED'}


class ResearchExportTests(unittest.TestCase):
    def ledger(self, key=('Product Details', 'RAM Memory Installed')):
        return {'schemaVersion': 1, 'productId': 'TEST-PC', 'configurationId': 'BASE-A',
                'facts': [{'id': 'sold', 'sheet': key[0], 'attribute': key[1],
                           'configurationId': 'BASE-A', 'fact_scope': 'SOLD_BASE',
                           'status': 'VERIFIED', 'final_value': 32, 'source': 'Synthetic fixture',
                           'applies_to_sold_configuration': True, 'transfer_basis': 'Test assembly record'}],
                'fields': [{'sheet': key[0], 'attribute': key[1], 'selected_fact': 'sold'}]}

    def test_sold_capacity_not_factory_capacity(self):
        data = self.ledger()
        factory = copy.deepcopy(data['facts'][0])
        factory.update(id='factory', fact_scope='FACTORY', final_value=16)
        data['facts'].append(factory)
        key = ('Product Details', 'RAM Memory Installed')
        self.assertEqual(r.project(data, [key], 'export')[key]['value'], 32)
        data['fields'][0]['selected_fact'] = 'factory'
        with self.assertRaisesRegex(ValueError, 'not factory'):
            r.project(data, [key], 'export')

    def test_platform_maximum_cannot_fill_installed_capacity(self):
        data = self.ledger()
        data['facts'][0].update(fact_scope='PLATFORM_CAPABILITY', final_value=64)
        with self.assertRaises(ValueError):
            r.project(data, [('Product Details', 'RAM Memory Installed')])
        key = ('Product Details', 'RAM Memory Maximum Size')
        data['facts'][0]['attribute'] = data['fields'][0]['attribute'] = key[1]
        self.assertEqual(r.project(data, [key], 'export')[key]['value'], 64)

    def test_fcc_candidate_stays_draft_and_cannot_be_promoted(self):
        key = ('Safety&Compliance', 'FCC ID')
        data = self.ledger(key)
        fact = data['facts'][0]
        fact.update(status='INPUT_UNVERIFIED', final_value=None, recommended_candidate='TEST-ID')
        draft = r.project(data, [key])
        self.assertIsNone(draft[key]['value'])
        self.assertEqual(draft[key]['value_candidate'], 'TEST-ID')
        self.assertIn('NOT CONFIRMED', draft[key]['value_candidate_source'])
        self.assertEqual(r.project(data, [key], separate_candidates=False)[key]['value'], 'TEST-ID')
        with self.assertRaisesRegex(ValueError, 'unresolved'):
            r.project(data, [key], 'export')
        fact.update(status='VERIFIED', final_value='TEST-ID', evidence_scope='RADIO_MODULE',
                    authorization_route='certification')
        with self.assertRaisesRegex(ValueError, 'candidate'):
            r.project(data, [key], 'export')
        del fact['recommended_candidate']
        fact['evidence'] = fcc_evidence()
        self.assertEqual(r.project(data, [key], 'export')[key]['value'], 'TEST-ID')

    def test_verified_fcc_needs_certification_and_scope(self):
        key = ('Safety&Compliance', 'FCC ID')
        data = self.ledger(key)
        data['facts'][0]['final_value'] = 'TEST-ID'
        with self.assertRaisesRegex(ValueError, 'certification'):
            r.project(data, [key], 'export')
        data['facts'][0].update(evidence_scope='HOST', authorization_route='certification', final_value='SDoC')
        with self.assertRaisesRegex(ValueError, 'identifier'):
            r.project(data, [key], 'export')

    def test_no_final_value_for_unresolved_and_no_formula(self):
        for changes in ({'status': 'TBD'}, {'final_value': '=1+1'}, {'final_value': None}, {'final_value': float('nan')}):
            data = self.ledger()
            data['facts'][0].update(changes)
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                r.validate_ledger(data)

    def test_mapping_is_complete_unique_and_configuration_bound(self):
        for change in ('duplicate', 'missing', 'wrong_attribute', 'wrong_config', 'missing_basis'):
            data = self.ledger()
            if change == 'duplicate': data['fields'].append(copy.deepcopy(data['fields'][0]))
            if change == 'missing': data['fields'][0]['selected_fact'] = 'unknown'
            if change == 'wrong_attribute': data['facts'][0]['attribute'] = 'other'
            if change == 'wrong_config': data['facts'][0]['configurationId'] = 'BASE-B'
            if change == 'missing_basis': data['facts'][0].pop('transfer_basis')
            with self.subTest(change=change), self.assertRaises(ValueError):
                r.project(data, [('Product Details', 'RAM Memory Installed')])
        with self.assertRaisesRegex(ValueError, 'exact current template'):
            r.project(self.ledger(), [])

    def test_unresolved_draft_allowed_and_na_is_blank(self):
        data = self.ledger()
        key = ('Product Details', 'RAM Memory Installed')
        data['facts'][0].update(status='TBD', final_value=None)
        self.assertIsNone(r.project(data, [key])[key]['value'])
        with self.assertRaises(ValueError): r.project(data, [key], 'export')
        data['facts'][0]['status'] = 'NOT_APPLICABLE'
        self.assertIsNone(r.project(data, [key], 'export')[key]['value'])

    def test_schema_rejects_legacy_and_reports_without_dumping_values(self):
        with self.assertRaises(ValueError) as error:
            r.validate_ledger([{'private': 'DO-NOT-PRINT'}])
        self.assertNotIn('DO-NOT-PRINT', str(error.exception))

    def full_fixture(self, root, name='review.xlsx', copy_present=True):
        book = openpyxl.load_workbook(r.ROOT/'assets/listing-workbook-template.xlsx')
        data = {'schemaVersion': 1, 'productId': 'TEST-PC', 'configurationId': 'BASE-A', 'facts': [], 'fields': []}
        data['serial_research'] = {
            'workflow_mode': 'FULL_LISTING_WORKFLOW', 'outcome': 'SERIAL_NOT_FOUND',
            'reason': 'Synthetic ERP has no linked unit serial', 'reviewed_date': '2026-10-14',
            'sources_attempted': ['Synthetic linked current and historical ERP records'],
            'units': [], 'resolved_fact_ids': [], 'factory_sold_differences': [],
            'unresolved_fields': ['Unit-specific factory BOM'], 'reconciliation': '',
            'next_action': 'Obtain linked unit evidence', 'owner': 'Synthetic researcher'}
        warranty = render('Dell')
        copy_values = {
            ('Product Details', 'Bullet Point'): warranty + '\n\nSynthetic bullet' * 4,
            ('Product Details', 'Product Description'): '<strong>Synthetic Product</strong><br/><br/>'
                '<strong>Warranty and Disclosure</strong><br/>' + warranty,
            ('Safety&Compliance', 'Warranty Description'): warranty,
            ('Product Details','Item Name'): 'MegaPC Custom Desktop Created Using Dell QCM1250 16GB RAM 256GB SSD Win 11 Pro',
            ('Product Details','Model Number'): 'QCM1250',
            ('Product Details','Model Name'): 'Pro Micro',
            ('Product Details','Manufacturer'): 'Dell',
            ('Product Details','CPU Model'): 'Core Ultra 7',
            ('Product Details','CPU Model Number'): '265T',
            ('Product Details','RAM Memory Installed'): 16,
            ('Product Details','RAM Memory Installed Size Unit'): 'GB',
            ('Product Details','Hard Disk Size'): 256,
            ('Product Details','Hard Disk Size Unit'): 'GB',
            ('Product Details','Operating System'): 'Windows 11 Pro',
        } if copy_present else {}
        for index, key in enumerate(r.workbook_fields(book)):
            fact = self.ledger(key)['facts'][0]
            fact.update(id=str(index), status='VERIFIED' if key in copy_values else 'NOT_APPLICABLE',
                        final_value=copy_values.get(key))
            if key == r.FCC_KEY:
                fact.update(evidence_scope='RADIO_MODULE', authorization_route='unresolved',
                            evidence={'lookup_outcome': 'FACTORY_NO_WLAN',
                                      'installation_basis': 'Synthetic factory no-WLAN BOM',
                                      'replacement_check': 'Synthetic sold assembly has no added radios',
                                      'radio_inventory_checked': True, 'installed_radio_modules': []})
            data['facts'].append(fact)
            data['fields'].append({'sheet': key[0], 'attribute': key[1], 'selected_fact': str(index)})
        projected = r.project(data, r.workbook_fields(book))
        for sheet in book:
            for col in range(2, sheet.max_column + 1):
                key = (sheet.title, sheet.cell(1, col).value)
                sheet.cell(3, col).value = copy_values.get(key)
                sheet.cell(4, col).value = 'VERIFIED' if key in copy_values else 'NOT_APPLICABLE'
                sheet.cell(5, col).value = projected[key]['source']
        path, ledger = root/name, root/'ledger.json'
        book.save(path)
        ledger.write_text(json.dumps(data), encoding='utf-8')
        (root/'product-facts.json').write_text(json.dumps({'identity': {'internalId':'TEST-PC','oemBrand':'Dell','oemModel':'QCM1250','cpu':'Intel Core Ultra 7 265T'},'selectedBaseConfiguration':{'id':'BASE-A','ramGB':16,'ssdGB':256}}), encoding='utf-8')
        return book, data, path, ledger

    def export_run(self, ledger, path, template=None, stage='LISTING_DRAFT'):
        # Every explicit synthetic export scenario supplies its current fixture audit.
        # Stale/missing-audit regressions exercise the production API directly.
        from reference_fixtures import write_audit
        write_audit(ledger)
        return r.run(template or r.ROOT/'assets/listing-workbook-template.xlsx', ledger, path,
                     'export', 'TEST-PC', 'BASE-A', stage)

    def test_cli_checks_workbook_tampering_and_target_identity(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            book, data, path, ledger = self.full_fixture(root)
            (root/'research-ledger.json').write_text(json.dumps(data), encoding='utf-8')
            from reference_fixtures import write_audit
            write_audit(ledger)
            command = [sys.executable, str(r.ROOT/'scripts/research_export.py'), '--ledger', str(ledger),
                       '--workbook', str(path), '--mode', 'export', '--product-id', 'TEST-PC',
                       '--stage', 'LISTING_DRAFT', '--configuration-id', 'BASE-A']
            passed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(passed.returncode, 0, passed.stderr)
            self.assertFalse(json.loads(passed.stdout)['amazonEligibilityAssessed'])
            wrong = command.copy(); wrong[-1] = 'BASE-B'
            self.assertNotEqual(subprocess.run(wrong, capture_output=True).returncode, 0)
            book['Product Details'].cell(3, 2).value = 'unreviewed edit'
            book.save(path); book.close()
            failed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
            self.assertNotEqual(failed.returncode, 0)
            self.assertIn('differs from scoped ledger', failed.stderr)

    def test_export_checks_copy_regardless_of_filename_and_requires_stage(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            for name in ('review.xlsx', 'RESEARCH_REVIEW.xlsx', 'TEST_LISTING_DRAFT.xlsx'):
                with self.subTest(name=name):
                    book, data, path, ledger = self.full_fixture(root, name, copy_present=False)
                    book.close()
                    with self.assertRaisesRegex(ValueError, 'Bullet 1'):
                        self.export_run(ledger, path)
                    with self.assertRaisesRegex(ValueError, 'explicit LISTING_DRAFT'):
                        self.export_run(ledger, path, stage='RESEARCH_REVIEW')
                    with self.assertRaisesRegex(ValueError, 'explicit LISTING_DRAFT'):
                        self.export_run(ledger, path, stage=None)
                    self.assertEqual(r.run(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, path,
                                          'draft', stage='RESEARCH_REVIEW')['result'], 'DRAFT_PROJECTION_PASS')
            (root/'product-facts.json').write_text('{}', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'PRODUCT_METADATA_REQUIRED.*oemBrand'):
                self.export_run(ledger, path)

    def test_export_binds_current_template_and_its_hash(self):
        import hashlib
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            book, data, path, ledger = self.full_fixture(root)
            book.close()
            canonical = r.ROOT/'assets/listing-workbook-template.xlsx'
            same = root/'same-template.xlsx'; same.write_bytes(canonical.read_bytes())
            result = self.export_run(ledger, path, same)
            self.assertEqual(result['templateSha256'], hashlib.sha256(canonical.read_bytes()).hexdigest())
            old = root/'old-template.xlsx'
            changed = openpyxl.load_workbook(canonical)
            changed['Product Details'].cell(2,2).value = 'Old definition'
            changed.save(old); changed.close()
            with self.assertRaisesRegex(ValueError, 'canonical template'):
                self.export_run(ledger, path, old)
            short = openpyxl.Workbook(); short.remove(short.active)
            for sheet_name in ('Product Details', 'Offer', 'Safety&Compliance'):
                short.create_sheet(sheet_name).append(['Field / Attribute', 'One field'])
            short.save(old); short.close()
            with self.assertRaisesRegex(ValueError, 'canonical template'):
                self.export_run(ledger, path, old)

    def test_workbook_type_changes_fail_but_numeric_round_trip_passes(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            book, data, path, ledger = self.full_fixture(root)
            fact = next(f for f in data['facts'] if f['attribute']=='RAM Memory Installed')
            fact.update(status='VERIFIED', final_value=16)
            column = next(c.column for c in book['Product Details'][1] if c.value=='RAM Memory Installed')
            book['Product Details'].cell(4,column).value = 'VERIFIED'
            ledger.write_text(json.dumps(data), encoding='utf-8')
            for value in (True, '16', False):
                book['Product Details'].cell(3,column).value = value; book.save(path)
                with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'differs'):
                    self.export_run(ledger, path)
            book['Product Details'].cell(3,column).value = 16.0; book.save(path)
            self.assertEqual(self.export_run(ledger, path)['result'], 'EXPORT_EVIDENCE_CHECK_PASS')
            fact['final_value'] = True; ledger.write_text(json.dumps(data), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'differs'):
                self.export_run(ledger, path)
            book.close()


if __name__ == '__main__':
    unittest.main()
