import importlib.util
import tempfile
import unittest
import sys
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).parents[1]))

spec=importlib.util.spec_from_file_location('workbooks',Path(__file__).parents[1]/'check-workbooks.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)

class WorkbookChecks(unittest.TestCase):
    def test_verified_fcc_id_cannot_be_blank_placeholder_or_number(self):
        for value in (None, '', 'TBD', 'N/A', 'SDoC', 'Yes', 'FCC', 12345):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'actual text identifier'):
                module.check_fcc_id_value(value, 'VERIFIED')
        module.check_fcc_id_value('TEST-MODULE-ID', 'VERIFIED')  # Text only, not grant validation.
        for status in ('TBD', 'SOURCE_UNAVAILABLE', 'NOT_APPLICABLE', 'CONFLICT'):
            module.check_fcc_id_value(None, status)

    def test_canonical_template_keeps_fcc_fields_and_research_statuses(self):
        module.check_template_inputs(Path(__file__).parents[2]/'assets/listing-workbook-template.xlsx')

    def test_current_workbook_discovery_includes_non_vl_ids_without_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            expected=[]
            for name in ('VL-123', 'VA-456', 'VD-744'):
                path=root/'product generated photo'/name/'research.xlsx'
                path.parent.mkdir(parents=True);path.touch();expected.append(path)
            for rel in ('history/product generated photo/VD-744/old.xlsx',
                        'product generated photo/VD-744/scratch/experiment.xlsx'):
                path=root/rel;path.parent.mkdir(parents=True);path.touch()
            self.assertEqual(module.current_workbooks(root),sorted(expected))

    def test_ui_prompts_are_optional_but_status_vocabulary_is_consistent(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'template.xlsx'
            book=openpyxl.load_workbook(Path(__file__).parents[2]/'assets/listing-workbook-template.xlsx')
            for sheet in book:
                for validation in sheet.data_validations.dataValidation:
                    validation.prompt=None;validation.showInputMessage=False
            book.save(path);module.check_template_inputs(path)
            book.worksheets[0].data_validations.dataValidation[0].formula1='"VERIFIED,NOT VERIFIED"'
            book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'Status dropdown'):
                module.check_template_inputs(path)

    def test_accepts_status_dispositions_but_rejects_structure_and_missing_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'template.xlsx'; book=openpyxl.Workbook();book.remove(book.active)
            for name in module.SHEETS:
                sheet=book.create_sheet(name)
                for r,label in enumerate(['Field / Attribute','Definition','Value','Status','Source'],1):sheet.cell(r,1,label)
                sheet.cell(1,2,'Installed RAM');sheet.cell(2,2,'Installed capacity only')
                sheet.cell(4,2,'TBD');sheet.cell(5,2,'SELLER_INPUT_REQUIRED: installed capacity')
            book.save(path);book.close();topology=module.template_topology(path)
            module.check_output(path,topology)
            book=openpyxl.load_workbook(path);book.worksheets[0].cell(1,2,'Maximum RAM');book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'headers/definitions'):module.check_output(path,topology)
            book=openpyxl.load_workbook(path);book.worksheets[0].cell(1,2,'Installed RAM');book.worksheets[0].cell(5,2).value=None;book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'Status/Source'):module.check_output(path,topology)
