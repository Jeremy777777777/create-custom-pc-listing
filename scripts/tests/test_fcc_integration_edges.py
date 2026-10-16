"""Regression cases for FCC evidence and workbook/summary handoff boundaries."""
import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import test_fcc_guide_writer as writer_fixtures
import test_research_export as export_fixtures
import research_export as r
import openpyxl
from openpyxl.styles import PatternFill


class FCCIntegrationEdgeTests(unittest.TestCase):
    def confirmed(self, root):
        data, fact, previous, ledger = writer_fixtures.FCCGuideWriterTests().fixture(
            root, 'VERIFIED', 'UNIT_CONFIRMED')
        fact['evidence'].update(fcc_record_id='TEST-MODULE-ID', fcc_record_model='SYNTHETIC-MODULE',
                                fcc_applicant='Synthetic applicant', fcc_grant_date='2026-01-01',
                                fcc_access_result='REVIEWED')
        return data, fact, previous, ledger

    def cli(self, ledger, out, *extra):
        args = ['research_export.py', '--ledger', str(ledger), '--workbook', str(out),
                '--write-workbook', '--stage', 'RESEARCH_REVIEW', *map(str, extra)]
        with patch.object(sys, 'argv', args), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            r.main()

    def keys(self, data):
        return [(field['sheet'], field['attribute']) for field in data['fields']]

    def test_existing_export_entrypoint_requires_new_fcc_contract(self):
        data = export_fixtures.ResearchExportTests().ledger(r.FCC_KEY)
        fact = data['facts'][0]
        fact.update(final_value='TEST-MODULE-ID', evidence_scope='RADIO_MODULE', authorization_route='certification')
        self.assertEqual(r.project(data, [r.FCC_KEY], 'draft')[r.FCC_KEY]['status'], 'VERIFIED')
        with self.assertRaisesRegex(ValueError, 'FCC.*lookup_outcome'):
            r.project(data, [r.FCC_KEY], 'export')

    def test_sdoc_cannot_hide_unknown_or_installed_radios(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, fact, previous, ledger = self.confirmed(Path(scratch))
            fact.update(status='NOT_APPLICABLE', final_value=None, evidence_scope='HOST', authorization_route='SDoC')
            fact['evidence'] = {'lookup_outcome': 'HOST_SDOC', 'sdoc_source': 'Synthetic matched OEM statement',
                                'installation_basis': 'Synthetic matched unit', 'replacement_check': 'No later additions'}
            with self.assertRaisesRegex(ValueError, 'radio'):
                r.project(data, self.keys(data), 'export')
            fact['evidence'].update(radio_inventory_checked=True, installed_radio_modules=['SYNTHETIC-RADIO'])
            with self.assertRaisesRegex(ValueError, 'radio'):
                r.project(data, self.keys(data), 'export')
            fact['evidence']['installed_radio_modules'] = []
            self.assertIsNone(r.project(data, self.keys(data), 'export')[r.FCC_KEY]['value'])

    def test_grant_match_access_and_host_module_scope_are_explicit(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, fact, previous, ledger = self.confirmed(Path(scratch))
            for key, wrong in [('fcc_record_id', 'OTHER-ID'), ('fcc_record_model', 'OTHER-MODULE'),
                               ('fcc_access_result', 'SOURCE_UNAVAILABLE')]:
                old = fact['evidence'][key]
                fact['evidence'][key] = wrong
                with self.subTest(key=key), self.assertRaises(ValueError):
                    r.validate_ledger(data)
                fact['evidence'][key] = old
            fact.update(evidence_scope='HOST')
            fact['evidence'].pop('module_model')
            fact['evidence'].update(host_model='SYNTHETIC-HOST', fcc_record_model='SYNTHETIC-HOST')
            self.assertEqual(r.project(data, self.keys(data), 'export')[r.FCC_KEY]['value'], 'TEST-MODULE-ID')

    def test_summary_conflict_is_detected_before_creating_excel(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.confirmed(root)
            ledger.write_text(json.dumps(data), encoding='utf-8')
            out = root/'new.xlsx'
            summary = out.with_suffix('.research-summary.json')
            summary.write_text('{"artifactRole":"UNRELATED"}', encoding='utf-8')
            before = summary.read_bytes()
            with self.assertRaises(SystemExit): self.cli(ledger, out)
            self.assertFalse(out.exists())
            self.assertEqual(summary.read_bytes(), before)

    def test_summary_must_stay_with_its_workbook_and_configuration(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.confirmed(root)
            ledger.write_text(json.dumps(data), encoding='utf-8')
            out = root/'new.xlsx'
            elsewhere = root/'elsewhere'; elsewhere.mkdir()
            with self.assertRaises(SystemExit):
                self.cli(ledger, out, '--summary-output', elsewhere/'new.research-summary.json')
            self.assertFalse(out.exists())
            summary = out.with_suffix('.research-summary.json')
            summary.write_text(json.dumps({'artifactRole': 'RESEARCH_COVERAGE_SUMMARY',
                                           'productId': 'TEST-PC', 'configurationId': 'OTHER',
                                           'workbookFile': out.name}), encoding='utf-8')
            with self.assertRaises(SystemExit): self.cli(ledger, out)
            self.assertFalse(out.exists())

    def test_summary_replace_failure_rolls_back_excel_and_preserves_old_sidecar(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.confirmed(root)
            ledger.write_text(json.dumps(data), encoding='utf-8')
            out = root/'new.xlsx'
            summary = out.with_suffix('.research-summary.json')
            summary.write_text(json.dumps({'artifactRole': 'RESEARCH_COVERAGE_SUMMARY',
                                           'productId': 'TEST-PC', 'configurationId': 'BASE-A',
                                           'workbookFile': out.name}), encoding='utf-8')
            before, prior = summary.read_bytes(), previous.read_bytes()
            with patch.object(Path, 'replace', side_effect=OSError('Synthetic summary disk failure')):
                with self.assertRaises(SystemExit): self.cli(ledger, out)
            self.assertFalse(out.exists())
            self.assertEqual(summary.read_bytes(), before)
            self.assertEqual(previous.read_bytes(), prior)
            self.assertFalse(list(root.glob('.*.tmp')))

    def test_canonical_stale_current_workbook_fails_before_new_output(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)/'product generated photo'/'TEST-PC'
            root.mkdir(parents=True)
            data, fact, previous, ledger = self.confirmed(root)
            ledger.write_text(json.dumps(data), encoding='utf-8')
            before = previous.read_bytes()
            out = root/'new.xlsx'
            with self.assertRaisesRegex(ValueError, 'Stale current workbook'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out)
            self.assertFalse(out.exists())
            self.assertEqual(previous.read_bytes(), before)

    def test_archived_previous_color_can_be_reused_only_from_a_matching_checked_bundle(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.confirmed(root)
            ledger.write_text(json.dumps(data), encoding='utf-8')
            old = openpyxl.load_workbook(previous)
            helper = writer_fixtures.FCCGuideWriterTests()
            helper.cell(old).value = 'TEST-MODULE-ID'
            helper.cell(old, 4).value = 'VERIFIED'
            helper.cell(old).fill = PatternFill('solid', fgColor='FFFF00')
            old.save(previous); old.close()
            bundle = root/'workbook-snapshots'/'v1'; bundle.mkdir(parents=True)
            paths = {'workbook': previous, 'ledger': ledger, 'template': r.ROOT/'assets/listing-workbook-template.xlsx'}
            records = {}
            for role, source in paths.items():
                destination = bundle/('template.xlsx' if role == 'template' else source.name)
                shutil.copyfile(source, destination)
                records[role] = {'path': destination.name, 'sha256': r.hashlib.sha256(destination.read_bytes()).hexdigest()}
            receipt = {'schemaVersion': 1, 'artifactRole': 'HISTORICAL_SNAPSHOT', 'exportEligible': False,
                       'productId': 'TEST-PC', 'configurationId': 'BASE-A', 'stage': 'RESEARCH_REVIEW', 'files': records}
            (bundle/'snapshot.json').write_text(json.dumps(receipt), encoding='utf-8')
            out = root/'new.xlsx'
            r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out,
                             previous_workbook=bundle/previous.name)
            book = openpyxl.load_workbook(out)
            self.assertEqual(helper.cell(book).fill.fgColor.rgb, '00FFFF00')
            book.close()
            receipt['configurationId'] = 'OTHER'
            (bundle/'snapshot.json').write_text(json.dumps(receipt), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'does not match'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, root/'wrong.xlsx',
                                 previous_workbook=bundle/previous.name)


if __name__ == '__main__':
    unittest.main()
