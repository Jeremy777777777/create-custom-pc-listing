import importlib.util
import unittest
from pathlib import Path
import openpyxl

spec=importlib.util.spec_from_file_location('warranty',Path(__file__).parents[1]/'warranty_policy.py')
w=importlib.util.module_from_spec(spec);spec.loader.exec_module(w)

class WarrantyTests(unittest.TestCase):
    def copy(self,brand):
        body=w.render(brand)
        return body, body+'\n\nOther verified benefit', '<strong>Computer</strong><br/><br/><strong>Warranty and Disclosure</strong><br/>'+body

    def test_same_policy_for_every_manufacturer(self):
        for brand in ['Lenovo','HP','Dell','ASUS']:
            body,bullets,description=self.copy(brand)
            self.assertEqual(w.validate_copy(brand,bullets,description,body)['fixedWording'],'PASS')
            self.assertFalse(w.validate_copy(brand,bullets,description,body)['oemApplicabilityOrPublicationReadinessInferred'])
            self.assertEqual(body,w.POLICY['paragraph'].replace('[OEM brand]',brand))

    def test_old_duration_omission_brand_and_heading_are_rejected(self):
        body,bullets,description=self.copy('Lenovo')
        for changed in [body.replace('1-year','6-month'),body.split(' The original seal')[0],body.replace('Lenovo','HP'),body+' Additional terms.']:
            with self.assertRaises(ValueError):w.validate_copy('Lenovo',changed,description,body)
            with self.assertRaises(ValueError):w.validate_copy('Lenovo',bullets,description,changed)
        for changed in [description.replace('Warranty and Disclosure','Warranty'),description+'\nMore copy',description.replace('1-year','2-year')]:
            with self.assertRaises(ValueError):w.validate_copy('Lenovo',bullets,changed,body)
        for brand in ['',None,'[OEM brand]','【制造商名】']:
            with self.assertRaises(ValueError):w.render(brand)

    def test_html_sections_preserve_copy_and_reject_broken_format(self):
        body, bullets, description = self.copy('ASUS')
        description = description.replace(
            '<strong>Warranty and Disclosure</strong>',
            '<strong>Display</strong><br/>OLED &amp; verified detail.<br/><br/>'
            '<strong>Warranty and Disclosure</strong>')
        self.assertEqual(w.validate_copy('ASUS', bullets, description, body)['fixedWording'], 'PASS')
        invalid = [
            description.replace('<strong>Display</strong>', 'Display'),
            description.replace('<br/><br/>', '<br/>', 1),
            description.replace('<br/>', '</br>'),
            description.replace('</strong>', '', 1),
            description + '<br/><br/>',
            '**Computer**\n**Warranty and Disclosure**' + chr(92) + '\n' + body,
        ]
        for candidate in invalid:
            with self.subTest(candidate=candidate):
                with self.assertRaises(ValueError):
                    w.validate_copy('ASUS', bullets, candidate, body)

    def test_actual_workbook_field_mapping(self):
        body,bullets,description=self.copy('Lenovo');book=openpyxl.Workbook();book.remove(book.active)
        for sheet,fields in [('Product Details',{'Bullet Point':bullets,'Product Description':description}),('Safety&Compliance',{'Warranty Description':body})]:
            ws=book.create_sheet(sheet)
            for column,(name,value) in enumerate(fields.items(),2):ws.cell(1,column,name);ws.cell(3,column,value)
        w.validate_workbook_copy(book,'Lenovo')
        book['Safety&Compliance'].cell(3,2,body.replace('1-year','6-month'))
        with self.assertRaises(ValueError):w.validate_workbook_copy(book,'Lenovo')
        book.close()
