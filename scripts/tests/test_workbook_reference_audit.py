"""Executable regressions for saved references and internal/external display scope."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import openpyxl
import workbook_reference_audit as c
import research_export as export
from reference_fixtures import write_audit
import test_fcc_guide_writer as writer_fixtures


class ReferenceContractTests(unittest.TestCase):
    def data(self):
        facts=[]
        for i,name in enumerate(sorted(c.PANEL_FIELDS)):
            facts.append({'id':str(i),'sheet':'Product Details','attribute':name,'fact_scope':'SOLD_BASE',
                'status':'VERIFIED','final_value':60 if name=='Display Refresh Rate in Hertz' else 'pixels' if name=='Display Resolution Maximum Unit' else '1920 x 1200',
                'source':'https://example.org/model','evidence':{'display_scope':'INTERNAL_PANEL'}})
        return {'productId':'TEST-PC','configurationId':'BASE-A','facts':facts,
                'fields':[{'selected_fact':f['id']} for f in facts]}

    def audit(self,d):
        return {'schemaVersion':1,'productId':'TEST-PC','configurationId':'BASE-A',
            'references':[{'url':'https://example.org/model','resolvedUrl':'https://example.org/model',
                'observedAtUtc':'2026-01-01T00:00:00Z','method':'BROWSER','outcome':'OPENED',
                'title':'SYNTHETIC model','evidence':'Synthetic panel1920x1200 at60Hz','limitations':'Synthetic test only'}],
            'fieldReviews':[{'factId':f['id'],'factSha256':c.digest(f),'primaryUrl':'https://example.org/model',
                'basis':'Synthetic matched model panel','supportScope':'SOLD_BASE'} for f in d['facts']]}

    def test_panel_4k_external_and_rate_type_are_rejected(self):
        for attr,change in [('Display Maximum Resolution',{'final_value':'3840 x 2160'}),
            ('Display Resolution Maximum',{'fact_scope':'PLATFORM_CAPABILITY'}),
            ('Native Resolution',{'evidence':{'display_scope':'EXTERNAL_OUTPUT'}}),
            ('Display Refresh Rate in Hertz',{'final_value':'60'}),
            ('Display Refresh Rate in Hertz',{'final_value':True}),
            ('Display Refresh Rate in Hertz',{'final_value':float('inf')}),
            ('Display Resolution Maximum Unit',{'final_value':'Hz'})]:
            d=self.data();next(f for f in d['facts'] if f['attribute']==attr).update(change)
            with self.subTest(attr=attr,change=change),self.assertRaises(ValueError):c.validate(d,self.audit(d))
        d=self.data();self.assertEqual(c.validate(d,self.audit(d))['state'],'WORKBOOK_REFERENCE_CONTRACT_PASS')

    def test_bad_access_or_snippet_cannot_be_opened_evidence(self):
        for change in [{'outcome':'NOT_FOUND'},{'outcome':'LOGIN_REQUIRED'},{'outcome':'ACCESS_BLOCKED'},
            {'outcome':'GONE'},{'outcome':'RENDER_UNAVAILABLE'},{'httpStatus':404},{'httpStatus':403},
            {'method':'HEAD'},{'observedAtUtc':'2999-01-01T00:00:00Z'},{'limitations':''}]:
            d=self.data();a=self.audit(d);a['references'][0].update(change)
            with self.subTest(change=change),self.assertRaises(ValueError):c.validate(d,a)

    def test_clean_targets_reject_notes_credentials_private_links(self):
        for url in ['http://example.org/model','https://example.org/model Reviewed today',
            'https://example.org/x|https://example.org/y','https://u:p@example.org/model',
            'https://example.org/?serial=PRIVATE','https://example.org/?X-Amz-Signature=secret']:
            with self.subTest(url=url),self.assertRaises(ValueError):c.clean_url(url)

    def test_hash_identity_and_field_coverage_are_bound(self):
        d=self.data()
        for mutation in ('identity','field','hash','missing','extra'):
            a=self.audit(d)
            if mutation=='identity':a['productId']='OTHER'
            if mutation=='field':a['fieldReviews'][0]['factSha256']='0'*64
            if mutation=='hash':a['ledgerSha256']='0'*64
            if mutation=='missing':a['fieldReviews'].pop()
            if mutation=='extra':a['fieldReviews'].append(copy.deepcopy(a['fieldReviews'][0]))
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):c.validate(d,a,b'actual' if mutation=='hash' else None)

    def test_saved_writer_native_links_and_missing_stale_audit(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);data,fact,old,ledger=writer_fixtures.FCCGuideWriterTests().fixture(root)
            out=root/'new.xlsx'
            result=export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,out)
            self.assertEqual(result['referenceAudit']['state'],'WORKBOOK_REFERENCE_CONTRACT_PASS')
            book=openpyxl.load_workbook(out);sh=book['Safety&Compliance'];col=next(cell.column for cell in sh[1] if cell.value=='FCC ID')
            self.assertEqual(sh.cell(5,col).hyperlink.target,fact['evidence']['fcc_official_url'])
            sh.cell(5,col).hyperlink='https://example.org/wrong';book.save(out);book.close()
            with self.assertRaisesRegex(ValueError,'native hyperlink'):c.check_paths(ledger,out,True)
            ledger.write_text(ledger.read_text()+'\n',encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'stale'):c.check_paths(ledger,required=True)
            (root/'reference-audit.json').unlink()
            with self.assertRaisesRegex(ValueError,'requires reference-audit'):export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,root/'blocked.xlsx')
            self.assertFalse((root/'blocked.xlsx').exists())
            self.assertEqual(c.check_paths(ledger)['state'],'LEGACY_REFERENCE_AUDIT_NOT_ASSESSED')

    def test_new_erp_policy_and_retired_active_sources(self):
        policy=json.loads((export.ROOT/'assets/erp-access-policy.json').read_text())
        self.assertEqual(policy['baseUrl'],'https://dev.mystore123.com')
        d=self.data();f=d['facts'][0];f['source']='https://'+policy['retiredHosts'][0]+'/products/1'
        a=self.audit(d);a['references'].append({**a['references'][0],'url':f['source'],'resolvedUrl':f['source']});a['fieldReviews'][0]['primaryUrl']=f['source']
        with self.assertRaisesRegex(ValueError,'Retired ERP'):c.validate(d,a)


if __name__=='__main__':unittest.main()
