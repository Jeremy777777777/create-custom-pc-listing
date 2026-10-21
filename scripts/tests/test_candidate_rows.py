"""Candidates are auditable draft advice and never leak into confirmed values."""
import copy
import sys
import unittest
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).parents[1]))
import research_export as r


class CandidateRowsTests(unittest.TestCase):
    key = ('Product Details', 'CPU Base Speed')

    def ledger(self, value=3.55):
        return {'schemaVersion':1,'productId':'TEST-PC','configurationId':'BASE',
                'facts':[{'id':'cpu','sheet':self.key[0],'attribute':self.key[1],
                          'configurationId':'BASE','fact_scope':'SOLD_BASE','status':'CONFLICT',
                          'final_value':None,'source':'Exact CPU suffix unresolved',
                          'value_candidate':value,'value_candidate_source':'https://example.test/official-cpu',
                          'evidence':{'candidate_rationale':'Two exact model sources favor this CPU',
                                      'candidate_missing_confirmation':'Unit CPU suffix conflicts with support page'}}],
                'fields':[{'sheet':self.key[0],'attribute':self.key[1],'selected_fact':'cpu'}]}

    def book(self, data, modern=True):
        b=openpyxl.Workbook();s=b.active;s.title=self.key[0]
        for row,label in enumerate(['Field / Attribute','Definition','Value','Status','Source']+
                                   (['Value Candidate','Value Candidate Source'] if modern else []),1):
            s.cell(row,1,label)
        s.cell(1,2,self.key[1]);s.cell(2,2,'Base frequency')
        p=r.project(data,[self.key])[self.key]
        for row,name in [(3,'value'),(4,'status'),(5,'source'),(6,'value_candidate'),(7,'value_candidate_source')]:
            if modern or row<6:s.cell(row,2,p[name])
        return b

    def test_candidate_provenance_and_no_promotion(self):
        d=self.ledger();p=r.project(d,[self.key])[self.key]
        self.assertIsNone(p['value']);self.assertEqual(p['status'],'CONFLICT')
        self.assertEqual(p['value_candidate'],3.55)
        for text in ['NOT CONFIRMED','https://example.test','Two exact','Unit CPU suffix']:
            self.assertIn(text,p['value_candidate_source'])
        with self.assertRaisesRegex(ValueError,'unresolved'):r.project(d,[self.key],'export')
        d['facts'][0].update(status='VERIFIED',final_value=3.55)
        with self.assertRaisesRegex(ValueError,'Candidate'):r.validate_ledger(d)

    def test_invalid_and_incomplete_metadata(self):
        for value in [None,'','  ','=1+1',float('inf'),float('nan')]:
            with self.subTest(value=value),self.assertRaises(ValueError):r.validate_ledger(self.ledger(value))
        for name in ['value_candidate','value_candidate_source']:
            d=self.ledger();del d['facts'][0][name]
            with self.assertRaises(ValueError):r.validate_ledger(d)
        for name in ['candidate_rationale','candidate_missing_confirmation']:
            d=self.ledger();del d['facts'][0]['evidence'][name]
            with self.assertRaisesRegex(ValueError,'requires evidence'):r.validate_ledger(d)
        for value in [0,False]:
            self.assertIs(r.project(self.ledger(value),[self.key])[self.key]['value_candidate'],value)

    def test_tampering_formula_and_missing_rows_fail(self):
        d=self.ledger();b=self.book(d);r.verify_workbook(b,d)
        for row,value in [(3,3.55),(6,3.2),(7,'different source'),(6,'=1+1'),(7,'=A1')]:
            altered=copy.deepcopy(b);altered.active.cell(row,2,value)
            with self.assertRaises(ValueError):r.verify_workbook(altered,d)
            altered.close()
        with self.assertRaisesRegex(ValueError,'seven-row'):r.verify_workbook(self.book(d,False),d)
        b.active.cell(7,1,'Changed label')
        with self.assertRaisesRegex(ValueError,'labels'):r.verify_workbook(b,d)
        b.close()

    def test_no_candidate_is_blank_and_legacy_still_checkable(self):
        d=self.ledger();f=d['facts'][0]
        del f['value_candidate'];del f['value_candidate_source']
        b=self.book(d,False);r.verify_workbook(b,d)
        p=r.project(d,[self.key])[self.key]
        self.assertIsNone(p['value_candidate']);self.assertIsNone(p['value_candidate_source'])
        with self.assertRaisesRegex(ValueError,'seven-row'):r.verify_workbook(b,d,'export')
        b.close()


if __name__=='__main__':unittest.main()
