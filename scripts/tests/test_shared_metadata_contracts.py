"""Cross-entrypoint regressions for metadata and historical source continuity."""
import copy
from contextlib import contextmanager
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import test_research_export as ef
import test_workflow_hardening as gf
import test_fact_projection as pf
import gallery_engine as g
from product_metadata import hardware_projection,listing_view,selected_facts,check_metadata_consistency
from research_history import check_research_ledger_history


class SharedMetadataTests(unittest.TestCase):
    @contextmanager
    def fixture(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);helper=ef.ResearchExportTests()
            book,data,workbook,ledger=helper.full_fixture(root)
            self.set_field(book,data,'CPU Model Manufacturer','Intel')
            book.save(workbook);book.close();ledger.write_text(json.dumps(data))
            product=json.loads((root/'product-facts.json').read_text())
            yield root,helper,data,workbook,ledger,product

    def set_field(self,book,data,attribute,value):
        fact=next(f for f in data['facts'] if f['sheet']=='Product Details' and f['attribute']==attribute)
        fact.update(status='VERIFIED',final_value=value)
        sheet=book['Product Details']
        col=next(c for c in range(2,sheet.max_column+1) if sheet.cell(1,c).value==attribute)
        sheet.cell(3,col).value=value;sheet.cell(4,col).value='VERIFIED'

    def test_explicit_raw_capacity_and_os_conflicts_fail_both_entrypoints(self):
        with self.fixture() as (root,helper,data,workbook,ledger,original):
            for structured in (False,True):
                for field,value in (('ram','32GB DDR5'),('internalStorage','512GB SSD'),('os','Windows 10 Home')):
                    product=copy.deepcopy(original)
                    if not structured:
                        product={'productId':'TEST-PC','identity':'Dell QCM1250',
                                 'cpu':'Intel Core Ultra 7 265T','ram':'16GB DDR5','internalStorage':'256GB SSD','os':'Windows 11 Pro'}
                    product[field]=value;(root/'product-facts.json').write_text(json.dumps(product))
                    with self.subTest(structured=structured,field=field):
                        with self.assertRaisesRegex(ValueError,'cache/research'):helper.export_run(ledger,workbook)
                        with self.assertRaisesRegex(ValueError,'cache/research'):hardware_projection(data,product)
            (root/'product-facts.json').write_text(json.dumps(original))
            self.assertEqual(helper.export_run(ledger,workbook)['result'],'EXPORT_EVIDENCE_CHECK_PASS')
            hardware_projection(data,original)

    def test_exact_cpu_cache_and_duplicate_cpu_records_cannot_disagree(self):
        with self.fixture() as (root,helper,data,workbook,ledger,original):
            for change in ('cpuModelNumber','top_cpu','model_only'):
                product=copy.deepcopy(original)
                if change in ('cpuModelNumber','model_only'):product['identity']['cpuModelNumber']='235T'
                else:product['cpu']='Intel Core Ultra 5 235T'
                if change=='model_only':product['identity'].pop('cpu')
                (root/'product-facts.json').write_text(json.dumps(product))
                with self.subTest(change=change):
                    with self.assertRaisesRegex(ValueError,'CPU'):helper.export_run(ledger,workbook)
                    with self.assertRaisesRegex(ValueError,'CPU'):hardware_projection(data,product)
            for value in ('265T','Intel(R) Core(TM) Ultra 7 265T'):
                product=copy.deepcopy(original);product['identity']['cpuModelNumber']=value
                (root/'product-facts.json').write_text(json.dumps(product))
                helper.export_run(ledger,workbook);hardware_projection(data,product)
            model_only=copy.deepcopy(data)
            manufacturer=next(f for f in model_only['facts'] if f['attribute']=='CPU Model Manufacturer')
            manufacturer.update(status='NOT_APPLICABLE',final_value=None)
            hardware_projection(model_only,{'identity':{'cpuModelNumber':'265T'}})

    def test_verified_alias_is_shared_and_undeclared_alias_still_fails(self):
        with self.fixture() as (root,helper,data,workbook,ledger,product):
            import openpyxl
            book=openpyxl.load_workbook(workbook)
            self.set_field(book,data,'CPU Model','Core Ultra 7 Series 2');book.save(workbook);book.close()
            ledger.write_text(json.dumps(data))
            product['identity']['cpuFieldAliases']={'CPU Model':['Core Ultra 7 Series 2']}
            (root/'product-facts.json').write_text(json.dumps(product))
            # A factual alias can be valid without being an observed Amazon option.
            with self.assertRaisesRegex(ValueError,'closest reference candidate'):
                helper.export_run(ledger,workbook)
            hardware_projection(data,product)
            del product['identity']['cpuFieldAliases']
            (root/'product-facts.json').write_text(json.dumps(product))
            with self.assertRaisesRegex(ValueError,'CPU'):helper.export_run(ledger,workbook)
            with self.assertRaisesRegex(ValueError,'CPU'):hardware_projection(data,product)

    def test_base_drive_medium_cannot_disagree_with_ssd_contract(self):
        with self.fixture() as (root,helper,data,workbook,ledger,product):
            import openpyxl
            for medium in ('HDD','Hard Disk Drive','eMMC','SSHD'):
                book=openpyxl.load_workbook(workbook);self.set_field(book,data,'Hard Disk Description',medium)
                book.save(workbook);book.close();ledger.write_text(json.dumps(data))
                with self.subTest(medium=medium):
                    with self.assertRaisesRegex(ValueError,'storage medium'):helper.export_run(ledger,workbook)
                    with self.assertRaisesRegex(ValueError,'storage medium'):hardware_projection(data,product)
            book=openpyxl.load_workbook(workbook);self.set_field(book,data,'Hard Disk Description','Solid State Drive')
            book.save(workbook);book.close();ledger.write_text(json.dumps(data))
            helper.export_run(ledger,workbook);hardware_projection(data,product)
            product['internalStorage']='256GB HDD'
            (root/'product-facts.json').write_text(json.dumps(product))
            with self.assertRaisesRegex(ValueError,'storage medium'):helper.export_run(ledger,workbook)
            with self.assertRaisesRegex(ValueError,'storage medium'):hardware_projection(data,product)
            product.pop('internalStorage')
            base=next(f for f in data['facts'] if f['attribute']=='Hard Disk Size')
            base['evidence']={'storage_medium':'INTERNAL_HDD'}
            ledger.write_text(json.dumps(data));(root/'product-facts.json').write_text(json.dumps(product))
            with self.assertRaisesRegex(ValueError,'storage medium'):helper.export_run(ledger,workbook)
            with self.assertRaisesRegex(ValueError,'storage medium'):hardware_projection(data,product)

    def test_factory_evidence_and_unresolved_research_are_not_promoted_or_blocked(self):
        with self.fixture() as (_,helper,data,workbook,ledger,product):
            original=copy.deepcopy(product)
            factory=copy.deepcopy(next(f for f in data['facts'] if f['attribute']=='RAM Memory Installed'))
            factory.update(id='factory-ram',fact_scope='FACTORY',final_value=32)
            data['facts'].append(factory)
            product['factoryConfiguration']={'ram':'32GB','os':'Windows 11 Home'}
            listing_view(product,data);hardware_projection(data,product)
            self.assertEqual(original['selectedBaseConfiguration'],product['selectedBaseConfiguration'])
            selected=selected_facts(data)
            selected[('Product Details','RAM Memory Installed')].update(status='TBD',final_value=None)
            check_metadata_consistency(selected,product)  # Draft may retain unresolved evidence.
            with self.assertRaisesRegex(ValueError,'unresolved'):
                check_metadata_consistency(selected,product,require_verified=True)


class ResearchContinuityTests(unittest.TestCase):
    def test_missing_git_only_skips_unversioned_scratch_checks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with patch('research_history.subprocess.run',side_effect=FileNotFoundError('git')):
                self.assertEqual(check_research_ledger_history(root),0)
                with self.assertRaisesRegex(ValueError,'Git is required'):
                    check_research_ledger_history(root,'base','head')
                (root/'.git').mkdir()
                with self.assertRaisesRegex(ValueError,'Git is required'):check_research_ledger_history(root)

    def test_real_git_deletion_downgrade_and_restoration_keep_historical_classification(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);p=root/'product generated photo/TEST';p.mkdir(parents=True)
            def git(*args):
                return subprocess.check_output(['git',*args],cwd=root,stderr=subprocess.DEVNULL).decode().strip()
            git('init');git('config','user.name','Test');git('config','user.email','test@localhost')
            fixture=gf.GalleryBoundaries()
            try:
                fixture.build(p);fixture.reviews(p,['PT08.png'])
                g.write(p/'delivery-status.json',{'galleryState':'LEGACY_QUARANTINED','currentQaPass':False})
                g.accept(p,slots=['PT08.png'])
                git('add','.');git('commit','-m','Pre-ledger synthetic review');before=git('rev-parse','HEAD')
                ledger=p/'research-ledger.json';scoped=pf.ResearchImageTests().ledger();g.write(ledger,scoped)
                git('add','.');git('commit','-m','Add scoped research');base=git('rev-parse','HEAD')
                self.assertEqual(g.gate(root,before,base)['historicalResearchUnboundPartialReviews'],1)
                reviewed={name:(p/name).read_bytes() for name in ('visual-review.json','semantic-review.json','partial-update-qa.json','delivery-status.json')}
                ledger.unlink()
                with self.assertRaisesRegex(ValueError,'ledger removed'):g.gate(root)
                with self.assertRaisesRegex(ValueError,'ledger removed'):g.review_inputs(p)
                with self.assertRaisesRegex(ValueError,'ledger removed'):g.accept(p,slots=['PT08.png'])
                workbook=p/'review.xlsx'
                shutil.copyfile(ef.r.ROOT/'assets/listing-workbook-template.xlsx',workbook)
                with self.assertRaisesRegex(ValueError,'ledger removed'):
                    ef.r.run(ef.r.ROOT/'assets/listing-workbook-template.xlsx',ledger,workbook,'draft')
                git('add','-A');git('commit','-m','Delete research')
                with self.assertRaisesRegex(ValueError,'ledger removed'):g.gate(root,base,'HEAD')
                g.write(ledger,[]);git('add','.');git('commit','-m','Downgrade research')
                with self.assertRaisesRegex(ValueError,'ledger downgraded'):g.gate(root,base,'HEAD')
                assets=root/'assets';assets.mkdir();shutil.copyfile(ef.r.ROOT/'assets/listing-workbook-template.xlsx',assets/'listing-workbook-template.xlsx')
                result=subprocess.run([sys.executable,str(ef.r.ROOT/'scripts/check-workbooks.py'),'--directory',str(root),'--base',base,'--head','HEAD'],capture_output=True,text=True)
                self.assertNotEqual(result.returncode,0);self.assertIn('ledger downgraded',result.stderr)
                g.write(ledger,scoped);git('add','.');git('commit','-m','Restore scoped research')
                summary=g.gate(root,base,'HEAD')
                self.assertEqual(summary['partialUpdatesValidated'],0)
                self.assertEqual(summary['historicalResearchUnboundPartialReviews'],1)
                self.assertEqual(reviewed,{name:(p/name).read_bytes() for name in reviewed})
            finally:fixture.doCleanups()

    def test_legacy_arrays_remain_supported_and_new_scoped_ledger_can_be_added(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);p=root/'product generated photo/TEST';p.mkdir(parents=True)
            def git(*args):return subprocess.check_output(['git',*args],cwd=root,stderr=subprocess.DEVNULL).decode().strip()
            git('init');git('config','user.name','Test');git('config','user.email','test@localhost')
            ledger=p/'research-ledger.json';g.write(ledger,[])
            git('add','.');git('commit','-m','Legacy');base=git('rev-parse','HEAD')
            g.write(ledger,[{'historical':'unverified'}])
            self.assertEqual(check_research_ledger_history(root),0)
            g.write(ledger,pf.ResearchImageTests().ledger());git('add','.');git('commit','-m','New scoped evidence')
            self.assertEqual(check_research_ledger_history(root,base,'HEAD'),0)
            self.assertEqual(check_research_ledger_history(root),1)


if __name__=='__main__':unittest.main()
