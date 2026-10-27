"""Synthetic schema tests; these inputs never claim a real browser lookup."""
import copy
import json
import sys
import tempfile
import subprocess
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parents[1]))
sys.path.insert(0, str(Path(__file__).parent))
import amazon_selection as a


class AmazonSelectionTests(unittest.TestCase):
    key = ('Product Details', 'RAM Type')
    sha = 'a'*64

    def inputs(self, value='LPDDR5', kind='CLOSED'):
        context = dict(marketplaceId='TEST', locale='en_US', productType='TEST_NOTEBOOK', itemTypeKeyword='test')
        control = {'sheet': self.key[0], 'attribute': self.key[1], 'path': 'system_ram_type-*-value',
                   'kind': kind, 'complete': True, 'evidenceSha256': 'b'*64,
                   'options': [{'code': 'ddr5', 'label': 'DDR5 RAM'}, {'code': 'ddr4', 'label': 'DDR4 RAM'}]}
        review = {'sheet': self.key[0], 'attribute': self.key[1], 'path': control['path'],
                  'decision': 'APPROXIMATE', 'referenceOptionCode': 'ddr5',
                  'reason': 'Same memory generation is the nearest available schema label',
                  'differences': 'LPDDR5 is low-power memory and is not ordinary DDR5 RAM',
                  'missingConfirmation': 'Amazon schema lacks the verified LPDDR5 type'}
        data = {'facts': [{'final_value': value, 'status': 'VERIFIED'}], 'amazon_selection': {
            'schemaVersion': 1, 'catalogSha256': self.sha, 'context': context, 'fields': [review],
            'catalogReference': 'https://example.test/synthetic-ui-options'}}
        projected = {self.key: {'value': value, 'status': 'VERIFIED', 'source': 'Official exact-unit research',
                               'value_candidate': None, 'value_candidate_source': None}}
        return data, projected, {'schemaVersion': 1, 'contexts': [{**context, 'controls': [control]}]}

    def test_no_match_always_requires_approximate_reference(self):
        d, p, c = self.inputs(); original = copy.deepcopy((d, p))
        out, report = a.apply(d, p, c, self.sha)
        self.assertEqual((d, p), original)
        self.assertEqual(out[self.key]['value_candidate'], 'DDR5 RAM')
        self.assertIsNone(out[self.key]['value'])
        self.assertEqual(out[self.key]['status'], 'INPUT_UNVERIFIED')
        for text in ['REFERENCE ONLY', 'NOT CONFIRMED', 'low-power']:
            self.assertIn(text, out[self.key]['value_candidate_source'])
        self.assertEqual(len(report['approximateReferenceFields']), 1)
        for name in ['referenceOptionCode', 'reason', 'differences', 'missingConfirmation']:
            altered = copy.deepcopy(d); del altered['amazon_selection']['fields'][0][name]
            with self.subTest(name=name), self.assertRaises(ValueError): a.apply(altered, p, c, self.sha)
        d['amazon_selection']['fields'][0]['decision'] = 'DIRECT'
        with self.assertRaisesRegex(ValueError, 'mandatory'): a.apply(d, p, c, self.sha)

    def test_approximate_is_never_exportable_and_never_invented_option(self):
        d, p, c = self.inputs()
        with self.assertRaisesRegex(ValueError, 'approximate'): a.apply(d, p, c, self.sha, 'export')
        d['amazon_selection']['fields'][0]['referenceOptionCode'] = 'lpddr5'
        with self.assertRaisesRegex(ValueError, 'observed allowed'): a.apply(d, p, c, self.sha)

    def test_no_fuzzy_match_but_reviewed_exact_alias(self):
        d, p, c = self.inputs('DDR5 SDRAM'); control = c['contexts'][0]['controls'][0]
        self.assertIsNone(a.match_option(control, 'DDR5 SDRAM'))
        control['aliases'] = {'ddr5 sdram': 'ddr5'}
        self.assertEqual(a.match_option(control, ' DDR5 SDRAM ')['code'], 'ddr5')
        d['amazon_selection']['fields'][0]['decision'] = 'DIRECT'
        out, _ = a.apply(d, p, c, self.sha, 'export')
        self.assertEqual(out[self.key]['value'], 'DDR5 RAM')

    def test_current_catalog_has_only_exact_reviewed_ram_and_unit_aliases(self):
        catalog, _=a.read_catalog()
        for context in catalog['contexts']:
            control=next(c for c in context['controls'] if c.get('attribute')=='RAM Type')
            self.assertEqual(a.match_option(control,'DDR5')['code'],'ddr5_ram')
            self.assertEqual(a.match_option(control,'DDR4')['code'],'ddr4_sdram')
            self.assertIsNone(a.match_option(control,'LPDDR5'))
            self.assertIsNone(a.match_option(control,'LPDDR5X'))
        context=next(c for c in catalog['contexts'] if c['productType']=='NOTEBOOK_COMPUTER')
        unit=next(c for c in context['controls'] if c.get('attribute')=='Lithium Battery Energy Content Unit')
        self.assertEqual(a.match_option(unit,'Wh')['code'],'watt_hours')
        self.assertIsNone(a.match_option(unit,'mAh at unspecified voltage'))

    def test_suggestions_remain_editable(self):
        d, p, c = self.inputs(kind='SUGGESTED')
        d['amazon_selection']['fields'][0]['decision'] = 'DIRECT'
        out, _ = a.apply(d, p, c, self.sha)
        self.assertEqual(out[self.key]['value'], 'LPDDR5')

    def test_account_specific_choices_are_never_reclassified_as_free_text(self):
        d, p, c = self.inputs('Unvalidated shipping template', kind='ACCOUNT_CLOSED')
        with self.assertRaisesRegex(ValueError, 'account review'):a.apply(d, p, c, self.sha)
        d['amazon_selection']['fields'][0].update(decision='UNOBSERVED', reason='Fresh account options needed')
        out, report=a.apply(d, p, c, self.sha)
        self.assertIsNone(out[self.key]['value'])
        self.assertEqual(report['state'], 'AMAZON_SELECTION_PARTIAL')
        with self.assertRaisesRegex(ValueError, 'account-specific'):a.apply(d, p, c, self.sha, 'export')

    def test_context_hash_review_and_unknown_fail_closed(self):
        d, p, c = self.inputs()
        with self.assertRaisesRegex(ValueError, 'stale'): a.apply(d, p, c, 'c'*64)
        altered = copy.deepcopy(d); altered['amazon_selection']['context']['productType'] = 'TEST_AIO'
        with self.assertRaisesRegex(ValueError, 'context'): a.apply(altered, p, c, self.sha)
        altered = copy.deepcopy(d); altered['amazon_selection']['fields'] = []
        with self.assertRaisesRegex(ValueError, 'each template field'): a.apply(altered, p, c, self.sha)
        c['contexts'][0]['controls'] = []
        d['amazon_selection']['fields'][0].update(decision='UNOBSERVED', reason='Conditional field absent')
        _, report = a.apply(d, p, c, self.sha)
        self.assertEqual(report['state'], 'AMAZON_SELECTION_PARTIAL')
        with self.assertRaisesRegex(ValueError, 'unobserved'): a.apply(d, p, c, self.sha, 'export')

    def test_repeatable_members_are_checked_individually(self):
        d, p, c = self.inputs('DDR5 RAM; DDR4 RAM'); control = c['contexts'][0]['controls'][0]
        control['repeatable'] = True
        r = d['amazon_selection']['fields'][0]
        r.update(decision='DIRECT', members={'value': ['DDR5 RAM', 'DDR4 RAM']})
        out, _ = a.apply(d, p, c, self.sha)
        self.assertEqual(out[self.key]['value'], 'DDR5 RAM; DDR4 RAM')
        r['members']['value'][1] = 'LPDDR5'
        with self.assertRaisesRegex(ValueError, 'exact original'): a.apply(d, p, c, self.sha)

    def test_gamut_approximation_never_relabels_original_percentage(self):
        d, p, c = self.inputs('DCI-P3')
        old = self.key; name = ('Product Details', 'Color Gamut Name'); amount = ('Product Details', 'Color Gamut Value')
        p[name] = p.pop(old)
        p[amount] = {'value': 99, 'status': 'VERIFIED', 'source': 'Synthetic 99 percent DCI-P3 evidence',
                     'value_candidate': None, 'value_candidate_source': None}
        control = c['contexts'][0]['controls'][0]
        control['attribute'] = name[1]; control['options'] = [{'code': 'srgb', 'label': 'sRGB'}]
        r = d['amazon_selection']['fields'][0]
        r.update(attribute=name[1], referenceOptionCode='srgb',
                 reason='Nearest available comparison gamut; low confidence',
                 differences='DCI-P3 and sRGB have different boundaries; percentages cannot be copied')
        d['amazon_selection']['fields'].append({'sheet': amount[0], 'attribute': amount[1],
            'decision': 'UNOBSERVED', 'reason': 'Synthetic numeric control not in fixture'})
        out, report = a.apply(d, p, c, self.sha)
        self.assertEqual(out[name]['value_candidate'], 'sRGB')
        self.assertIsNone(out[amount]['value'])
        self.assertIsNone(out[amount]['value_candidate'])
        self.assertIn('99', out[amount]['source'])
        self.assertEqual(p[amount]['value'], 99)
        self.assertIn(list(amount), report['pairedBasisUnresolvedFields'])

    def test_real_writer_saved_bytes_preserve_fact_and_reject_candidate_tampering(self):
        import openpyxl
        import research_export as export
        from test_fcc_guide_writer import FCCGuideWriterTests
        from reference_fixtures import write_audit
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch)
            data, _, previous, ledger=FCCGuideWriterTests().fixture(root)
            fact=next(f for f in data['facts'] if (f['sheet'],f['attribute'])==self.key)
            fact.update(status='VERIFIED',final_value='LPDDR5')
            original_fact=copy.deepcopy(fact)
            catalog, sha=a.read_catalog()
            context=a.context_for(catalog,data['amazon_selection']['context'])
            control=next(c for c in context['controls'] if (c['sheet'],c['attribute'])==self.key)
            option=next(o for o in control['options'] if o['label']=='DDR5 RAM')
            review=next(f for f in data['amazon_selection']['fields'] if (f['sheet'],f['attribute'])==self.key)
            review.update(decision='APPROXIMATE',referenceOptionCode=option['code'],
                reason='Synthetic nearest same-generation label',differences='LPDDR5 is not ordinary DDR5',
                missingConfirmation='Matching exact Amazon type unavailable in synthetic case')
            data['amazon_selection']['catalogReference']='https://example.test/synthetic-ui-options'
            ledger.write_text(json.dumps(data),encoding='utf-8');write_audit(ledger)
            output=root/'new.xlsx'
            export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx',ledger,output)
            book=openpyxl.load_workbook(output)
            sheet=book[self.key[0]];col=next(c.column for c in sheet[1] if c.value==self.key[1])
            self.assertIsNone(sheet.cell(3,col).value)
            self.assertEqual(sheet.cell(6,col).value,'DDR5 RAM')
            self.assertEqual(sheet.cell(7,col).hyperlink.target,data['amazon_selection']['catalogReference'])
            self.assertEqual(fact,original_fact)
            export.verify_workbook(book,data)
            sheet.cell(6,col).value='LPDDR5'
            with self.assertRaisesRegex(ValueError,'differs'):export.verify_workbook(book,data)
            keys=export.workbook_fields(book)
            book.close()
            # The existing unresolved FCC fact may independently block first.
            with self.assertRaisesRegex(ValueError,'blocked'):export.project(data,keys,'export')

    def test_compiler_rejects_incomplete_lists_before_replacing_previous_catalog(self):
        with tempfile.TemporaryDirectory() as scratch:
            root=Path(scratch);output=root/'existing-catalog.json';output.write_text('preserved old bytes')
            capture={'schemaVersion':1,'marketplaceId':'TEST','locale':'en_US','productType':'TEST_TYPE',
                'itemTypeKeyword':'test','observedAtUtc':'2026-10-27T00:00:00Z','controls':[
                    {'tag':'KAT-DROPDOWN','name':'system_ram_type-0-value','label':'RAM Type',
                     'optionCount':2,'options':[{'value':'ddr5_ram','label':'DDR5 RAM'}]}]}
            (root/'seller-notebook-product-details.json').write_text(json.dumps(capture))
            command=[sys.executable,str(a.ROOT/'scripts/compile_amazon_selections.py'),
                     '--captures-dir',str(root),'--output',str(output)]
            failed=subprocess.run(command,capture_output=True,text=True)
            self.assertNotEqual(failed.returncode,0)
            self.assertEqual(output.read_text(),'preserved old bytes')
            self.assertEqual(len(list(root.glob('tmp*.json'))),0)
            capture['controls'][0]['optionCount']=1
            (root/'seller-notebook-product-details.json').write_text(json.dumps(capture))
            passed=subprocess.run(command,capture_output=True,text=True)
            self.assertEqual(passed.returncode,0,passed.stderr)
            compiled,_=a.read_catalog(output)
            self.assertEqual(len(compiled['contexts']),1)


if __name__ == '__main__':
    unittest.main()
