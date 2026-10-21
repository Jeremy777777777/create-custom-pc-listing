"""End-to-end synthetic Excel creation and guide/export boundary checks."""
import json
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
import test_research_export as fixtures
import research_export as r


class FCCGuideWriterTests(unittest.TestCase):
    def fixture(self, root, status='INPUT_UNVERIFIED', outcome='SUPPORTED_OPTION_SELECTED'):
        book, data, previous, ledger = fixtures.ResearchExportTests().full_fixture(root)
        fact = next(f for f in data['facts'] if (f['sheet'], f['attribute']) == r.FCC_KEY)
        fact.update(status=status, final_value='TEST-MODULE-ID' if status == 'VERIFIED' else None,
                    evidence_scope='RADIO_MODULE', authorization_route='certification',
                    evidence={**fixtures.fcc_evidence('TEST-MODULE-ID'), 'lookup_outcome': outcome,
                              'radio_inventory_checked': True, 'installed_radio_modules': [],
                              'module_model': 'SYNTHETIC-MODULE',
                              'installation_basis': 'Synthetic unit label',
                              'replacement_check': 'Synthetic assembly record: no replacement',
                              'fcc_official_url': 'https://apps.fcc.gov/oetcf/eas/reports/GenericSearch.cfm',
                              'reviewed_date': '2026-10-16'})
        if status == 'INPUT_UNVERIFIED':
            fact['recommended_candidate'] = 'TEST-MODULE-ID'
        ledger.write_text(json.dumps(data), encoding='utf-8')
        book.close()
        return data, fact, previous, ledger

    def cell(self, book, row=3):
        sheet = book['Safety&Compliance']
        col = next(c.column for c in sheet[1] if c.value == 'FCC ID')
        return sheet.cell(row, col)

    def test_selected_option_uses_candidate_rows_and_blocks_export(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.fixture(root)
            output = root/'new.xlsx'
            template_hash = r.hashlib.sha256((r.ROOT/'assets/listing-workbook-template.xlsx').read_bytes()).hexdigest()
            result = r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, output)
            self.assertEqual(result['result'], 'DRAFT_PROJECTION_PASS')
            book = openpyxl.load_workbook(output)
            template = openpyxl.load_workbook(r.ROOT/'assets/listing-workbook-template.xlsx')
            self.assertEqual(book.sheetnames, template.sheetnames)
            for sheet in book:
                for row in (1, 2):
                    self.assertEqual([c.value for c in sheet[row]], [c.value for c in template[sheet.title][row]])
                self.assertEqual(str(sheet.data_validations), str(template[sheet.title].data_validations))
            self.assertIsNone(self.cell(book).value)
            self.assertEqual(self.cell(book, 6).value, 'TEST-MODULE-ID')
            self.assertEqual(self.cell(book, 6).number_format, '@')
            self.assertEqual(self.cell(book, 6).fill.fgColor.rgb, '00FFF2CC')
            self.assertEqual(self.cell(book, 4).value, 'INPUT_UNVERIFIED')
            self.assertIn('USER_SELECTED_SUPPORTED_OPTION', self.cell(book, 7).value)
            self.assertNotIn('MOST_LIKELY', self.cell(book, 7).value)
            book.close(); template.close()
            self.assertEqual(template_hash, r.hashlib.sha256((r.ROOT/'assets/listing-workbook-template.xlsx').read_bytes()).hexdigest())
            blocked = root/'blocked.xlsx'
            with self.assertRaisesRegex(ValueError, 'unresolved'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, blocked,
                                 'export', 'TEST-PC', 'BASE-A', 'LISTING_DRAFT')
            self.assertFalse(blocked.exists())
            self.assertEqual(list(root.glob('.new-*.xlsx')), [])

    def test_pending_and_no_wlan_are_blank_red_with_distinct_statuses(self):
        for status, outcome in [('TBD', 'PENDING'), ('NOT_APPLICABLE', 'FACTORY_NO_WLAN')]:
            with self.subTest(status=status), tempfile.TemporaryDirectory() as scratch:
                root = Path(scratch)
                data, fact, previous, ledger = self.fixture(root, status, outcome)
                out = root/'new.xlsx'
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out)
                book = openpyxl.load_workbook(out)
                self.assertIsNone(self.cell(book).value)
                self.assertEqual(self.cell(book).fill.fgColor.rgb, '00FFC7CE')
                self.assertEqual(self.cell(book, 4).value, status)
                self.assertIn(outcome, self.cell(book, 5).value)
                book.close()

    def test_confirmed_export_and_historical_color_preservation(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.fixture(root, 'VERIFIED', 'UNIT_CONFIRMED')
            old = openpyxl.load_workbook(previous)
            self.cell(old).value = 'TEST-MODULE-ID'
            self.cell(old, 4).value = 'VERIFIED'
            self.cell(old).fill = PatternFill('solid', fgColor='FFFF00')
            old.save(previous); old.close()
            out = root/'new.xlsx'
            result = r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out,
                                     'export', 'TEST-PC', 'BASE-A', 'LISTING_DRAFT', previous)
            self.assertEqual(result['result'], 'EXPORT_EVIDENCE_CHECK_PASS')
            book = openpyxl.load_workbook(out)
            self.assertEqual(self.cell(book).fill.fgColor.rgb, '00FFFF00')
            self.assertIn('https://apps.fcc.gov/', self.cell(book, 5).value)
            book.close()
            before = out.read_bytes()
            with self.assertRaisesRegex(ValueError, 'already exists'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out)
            self.assertEqual(before, out.read_bytes())

    def test_false_confirmation_and_unchecked_no_radio_rejected(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, fact, previous, ledger = self.fixture(Path(scratch), 'VERIFIED', 'UNIT_CONFIRMED')
            for outcome in ('SUPPORTED_OPTION_SELECTED', 'MODEL_MODULE_MATCH', 'PENDING', 'FACTORY_NO_WLAN'):
                fact['evidence']['lookup_outcome'] = outcome
                with self.subTest(outcome=outcome), self.assertRaises(ValueError):
                    r.validate_ledger(data)
            fact['evidence']['lookup_outcome'] = 'UNIT_CONFIRMED'
            for key in ('installation_basis', 'replacement_check', 'fcc_official_url', 'reviewed_date', 'module_model'):
                old = fact['evidence'].pop(key)
                with self.subTest(key=key), self.assertRaises(ValueError):
                    r.validate_ledger(data)
                fact['evidence'][key] = old
            fact['evidence']['fcc_official_url'] = 'https://fcc.gov.example.com/grant'
            with self.assertRaises(ValueError): r.validate_ledger(data)
            fact.update(status='NOT_APPLICABLE', final_value=None)
            fact['evidence'].update(lookup_outcome='FACTORY_NO_WLAN', replacement_check='')
            with self.assertRaisesRegex(ValueError, 'later-addition'): r.validate_ledger(data)

    def test_cli_writes_bound_summary_and_rejects_bad_target_without_output(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.fixture(root)
            out = root/'new.xlsx'
            command = [sys.executable, str(r.ROOT/'scripts/research_export.py'), '--ledger', str(ledger),
                       '--workbook', str(out), '--write-workbook', '--stage', 'RESEARCH_REVIEW']
            completed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(completed.returncode, 0, completed.stderr)
            summary = json.loads(out.with_suffix('.research-summary.json').read_text(encoding='utf-8'))
            self.assertEqual(summary['workbookFile'], out.name)
            self.assertEqual(summary['workbookSha256'], r.hashlib.sha256(out.read_bytes()).hexdigest())
            blocked = root/'wrong.xlsx'
            with self.assertRaisesRegex(ValueError, 'another target'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, blocked, product_id='OTHER')
            self.assertFalse(blocked.exists())

    def test_failed_output_write_removes_partial_file_and_keeps_prior_bytes(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            data, fact, previous, ledger = self.fixture(root)
            prior = previous.read_bytes()
            out = root/'failed.xlsx'
            original_open = Path.open

            class FailedWriter:
                def __enter__(self):
                    self.handle = original_open(out, 'xb')
                    return self

                def write(self, payload):
                    self.handle.write(payload[:20])
                    raise OSError('Synthetic disk failure')

                def __exit__(self, *args):
                    self.handle.close()

            def intercept(path, *args, **kwargs):
                if path == out and args == ('xb',):
                    return FailedWriter()
                return original_open(path, *args, **kwargs)

            with patch.object(Path, 'open', intercept), self.assertRaisesRegex(OSError, 'Synthetic disk failure'):
                r.write_workbook(r.ROOT/'assets/listing-workbook-template.xlsx', ledger, out)
            self.assertFalse(out.exists())
            self.assertEqual(previous.read_bytes(), prior)
            self.assertEqual(list(root.glob('.failed-*.xlsx')), [])


if __name__ == '__main__':
    unittest.main()
