"""Synthetic regression cases, not research evidence or outputs for a real SKU."""
import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import customization_policy as policy
import research_export as export
from listing_content import validate_listing
from product_metadata import listing_view
from reference_fixtures import write_audit
import test_research_export as fixtures


def configure(data, product, replaceable=True, memory_type='DDR5'):
    """Supply conspicuously synthetic inspected-hardware evidence, then compose copy."""
    ram = next(f for f in data['facts'] if f['attribute'] == 'RAM Memory Installed')
    ram.setdefault('evidence', {})['customization_ram'] = {
        'status': 'VERIFIED', 'replaceable': replaceable, 'memoryType': memory_type,
        'replaceabilitySource': 'https://example.test/synthetic-service-manual#memory',
        'memoryTypeSource': 'https://example.test/synthetic-parts#memory', 'reviewedDate': '2026-10-29'}
    technology = next(f for f in data['facts'] if f['attribute'] == 'RAM Memory Technology')
    technology.update(status='VERIFIED', final_value=memory_type)
    view = listing_view(product, data)[0]
    texts = policy.copy_texts(view)
    for fact in data['facts']:
        if fact['attribute'] == 'Item Name':
            fact['final_value'] = ('MegaPC Custom Desktop, Created Using Dell QCM1250, '
                                   + texts['titleFragment']+', Win 11 Pro')
        elif fact['attribute'] == 'Bullet Point':
            paragraphs = fact['final_value'].split('\n\n')
            paragraphs[3] = texts['bullet']
            fact['final_value'] = '\n\n'.join(paragraphs)
        elif fact['attribute'] == 'Product Description':
            fact['final_value'] = fact['final_value'].replace(
                '<strong>Warranty and Disclosure</strong>',
                texts['descriptionSection']+'<br/><br/><strong>Warranty and Disclosure</strong>')
    return view


class CustomizationTests(unittest.TestCase):
    def fixture(self, root, replaceable=True, memory_type='DDR5'):
        book, data, old, ledger = fixtures.ResearchExportTests().full_fixture(root)
        book.close()
        product = json.loads((root/'product-facts.json').read_text())
        view = configure(data, product, replaceable, memory_type)
        return data, product, view, ledger

    def projection(self, data):
        return export.project(data, [(m['sheet'], m['attribute']) for m in data['fields']])

    def test_replaceable_ddr4_and_ddr5_have_full_choices_in_all_fields(self):
        for memory_type in ('DDR4', 'DDR5'):
            with self.subTest(memory_type=memory_type), tempfile.TemporaryDirectory() as scratch:
                data, product, view, _ = self.fixture(Path(scratch), True, memory_type)
                projected = self.projection(data)
                self.assertEqual(validate_listing(projected, view)['content'], 'PASS')
                for field in ('Item Name', 'Bullet Point', 'Product Description'):
                    text = projected[('Product Details', field)]['value']
                    self.assertIn(policy.RAM_OPTIONS+' '+memory_type, text)
                    self.assertIn(policy.SSD_OPTIONS, text)
                self.assertLessEqual(len(projected[('Product Details', 'Item Name')]['value']), 200)
                self.assertEqual(product['selectedBaseConfiguration']['ramGB'], 16)
                self.assertEqual(projected[('Product Details', 'Hard Disk Size')]['value'], 256)

    def test_fixed_lpddr_generations_keep_capacity_and_no_ram_choices(self):
        for memory_type in ('LPDDR4X', 'LPDDR5', 'LPDDR5X'):
            with self.subTest(memory_type=memory_type), tempfile.TemporaryDirectory() as scratch:
                data, _, view, _ = self.fixture(Path(scratch), False, memory_type)
                projected = self.projection(data)
                validate_listing(projected, view)
                for field in ('Item Name', 'Bullet Point', 'Product Description'):
                    text = projected[('Product Details', field)]['value']
                    self.assertIn('16GB '+memory_type+' RAM (Fixed)', text)
                    self.assertNotIn(policy.RAM_OPTIONS, text)
                    self.assertIn(policy.SSD_OPTIONS, text)

    def test_platform_maximum_does_not_filter_seller_ram_options(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, _, view, _ = self.fixture(Path(scratch))
            projected = self.projection(data)
            projected[('Product Details', 'RAM Memory Maximum Size')].update(value=32, status='VERIFIED')
            projected[('Product Details', 'RAM Memory Maximum Size Unit')].update(value='GB', status='VERIFIED')
            validate_listing(projected, view)
            self.assertIn('64GB', policy.copy_texts(view)['titleFragment'])
            self.assertEqual(projected[('Product Details', 'RAM Memory Maximum Size')]['value'], 32)

    def test_missing_unknown_conflict_and_incomplete_evidence_do_not_guess_branch(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, product, _, _ = self.fixture(Path(scratch))
            ram = next(f for f in data['facts'] if f['attribute'] == 'RAM Memory Installed')
            for status in ('TBD', 'CONFLICT', 'SOURCE_UNAVAILABLE'):
                ram['evidence']['customization_ram']['status'] = status
                with self.subTest(status=status), self.assertRaisesRegex(ValueError, 'unresolved'):
                    policy.research(product, data, required=True)
            ram['evidence'].pop('customization_ram')
            self.assertIsNone(policy.research(product, data))
            with self.assertRaisesRegex(ValueError, 'requires researched'):
                policy.research(product, data, required=True)

    def test_type_source_replaceability_and_cache_mismatches_fail(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, product, view, _ = self.fixture(Path(scratch))
            for changes in ({'memoryType': 'DDR'}, {'replaceable': 'yes'}, {'memoryTypeSource': ''},
                            {'reviewedDate': '2099-01-01'}, {'memoryType': 'DDR4'}):
                changed = copy.deepcopy(data)
                next(f for f in changed['facts'] if f['attribute'] == 'RAM Memory Installed')['evidence']['customization_ram'].update(changes)
                with self.subTest(changes=changes), self.assertRaises(ValueError):
                    policy.research(product, changed, required=True)
            cached = copy.deepcopy(product)
            cached['customization'] = {**view['customization'], 'memoryType': 'DDR4'}
            with self.assertRaisesRegex(ValueError, 'cache disagrees'):
                listing_view(cached, data)

    def test_wrong_generation_unsupported_choices_and_wrong_installed_claim_fail(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, _, view, _ = self.fixture(Path(scratch))
            for field in ('Item Name', 'Bullet Point', 'Product Description'):
                for transform in (lambda s: s.replace('DDR5', 'DDR4'),
                                  lambda s: s.replace('4TB', '8TB'),
                                  lambda s: s+'; 128GB RAM',
                                  lambda s: s+'; 2TB HDD',
                                  lambda s: s+'; 16GB LPDDR5 RAM'):
                    changed = self.projection(data)
                    changed[('Product Details', field)]['value'] = transform(changed[('Product Details', field)]['value'])
                    with self.subTest(field=field), self.assertRaises(ValueError):
                        validate_listing(changed, view)

    def test_fixed_ram_cannot_advertise_selectable_ram(self):
        with tempfile.TemporaryDirectory() as scratch:
            data, _, view, _ = self.fixture(Path(scratch), False, 'LPDDR5X')
            changed = self.projection(data)
            changed[('Product Details', 'Bullet Point')]['value'] += '; '+policy.RAM_OPTIONS+' LPDDR5X RAM Options'
            with self.assertRaises(ValueError):
                validate_listing(changed, view)

    def test_saved_workbooks_and_real_cli_use_researched_branch_and_readback(self):
        for replaceable, memory_type in ((True, 'DDR5'), (False, 'LPDDR5X')):
            with self.subTest(memory_type=memory_type), tempfile.TemporaryDirectory() as scratch:
                root = Path(scratch)
                data, product, view, ledger = self.fixture(root, replaceable, memory_type)
                for fact in data['facts']:
                    fact.setdefault('evidence', {})['research_trace'] = {
                        'version': 1, 'field_class': 'DERIVED_COPY' if fact['attribute'] in ('Item Name', 'Bullet Point', 'Product Description') else 'TECHNICAL',
                        'reason': 'SYNTHETIC TEST ONLY; tests recorded coverage, not actual research',
                        'basis': 'Synthetic known fixture inputs and seller policy',
                        'start': 'REUSE_VERIFIED', 'start_basis': 'Synthetic fixture input, no real OEM query',
                        'outcome': 'CONFIRMED' if fact['status'] == 'VERIFIED' else 'NOT_APPLICABLE',
                        'routes': [{'route': 'OEM_MODEL', 'query': 'synthetic exact model',
                                    'reference': 'https://example.test/synthetic', 'reviewed_date': '2026-10-29',
                                    'result': 'REVIEWED', 'detail': 'Synthetic test fixture only'}]}
                ledger.write_text(json.dumps(data), encoding='utf-8')
                write_audit(ledger)
                output = root/'TEST_LISTING_DRAFT.xlsx'
                command = [sys.executable, str(export.ROOT/'scripts/research_export.py'), '--ledger', str(ledger),
                           '--workbook', str(output), '--write-workbook', '--stage', 'LISTING_DRAFT']
                result = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
                self.assertEqual(result.returncode, 0, result.stderr)
                checked = export.run(export.ROOT/'assets/listing-workbook-template.xlsx', ledger, output,
                                     'draft', stage='LISTING_DRAFT', check_copy=True)
                self.assertEqual(checked['result'], 'DRAFT_PROJECTION_PASS')
                saved = openpyxl.load_workbook(output)
                values = {saved['Product Details'].cell(1, c).value: saved['Product Details'].cell(3, c).value
                          for c in range(2, saved['Product Details'].max_column+1)}
                policy.validate_copy(values, view)
                self.assertEqual(values['RAM Memory Installed'], 16)
                self.assertEqual(values['Hard Disk Size'], 256)
                self.assertEqual(len(values['Bullet Point'].split('\n\n')), 5)
                self.assertIn('<strong>Memory and Storage Customization</strong><br/>', values['Product Description'])
                self.assertTrue(values['Product Description'].endswith(values['Bullet Point'].split('\n\n')[0]))
                saved.close()

    def test_new_listing_writer_rejects_missing_ram_research_without_output(self):
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            book, _, _, ledger = fixtures.ResearchExportTests().full_fixture(root)
            book.close()
            output = root/'new.xlsx'
            with self.assertRaisesRegex(ValueError, 'requires researched'):
                export.write_workbook(export.ROOT/'assets/listing-workbook-template.xlsx', ledger, output, stage='LISTING_DRAFT')
            self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main()
