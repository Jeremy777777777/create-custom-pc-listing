"""Validate recorded browser evidence and native workbook links, not source truth.

This checker never logs in, performs research, or manufactures an observation.
New generation/export/delivery must supply the actual reference-audit.json.
"""
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
from urllib.parse import parse_qsl, unquote, urlsplit

URL_PATTERN = re.compile(r'https?://[^\s<>";|]+')
RESOLUTION_FIELDS = {'Display Maximum Resolution', 'Display Resolution Maximum',
                     'Resolution', 'Native Resolution'}
PANEL_FIELDS = RESOLUTION_FIELDS | {'Display Refresh Rate in Hertz', 'Display Resolution Maximum Unit'}
ERP_POLICY = Path(__file__).resolve().parents[1]/'assets/erp-access-policy.json'


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':')).encode('utf-8')).hexdigest()


def urls(text):
    return list(dict.fromkeys(m.group(0).rstrip('.,)') for m in URL_PATTERN.finditer(str(text or ''))))


def clean_url(value):
    if not isinstance(value, str) or not value or re.search(r'[\s<>"|]', value):
        raise ValueError('Reference target must be one clean HTTPS URL')
    u = urlsplit(value)
    if u.scheme != 'https' or not u.hostname or u.username or u.password:
        raise ValueError('Reference target requires HTTPS without credentials')
    decoded = unquote(value)
    if re.search(r'[\r\n]|\s(?:Reviewed|Sources?|Rationale|Missing)\b|https?://.*https?://', decoded):
        raise ValueError('Reference URL contains concatenated notes or another URL')
    sensitive = {'serial', 'serialnumber', 'serial_number', 'servicetag', 'service_tag',
                 'token', 'access_token', 'refresh_token', 'signature', 'sig', 'code'}
    if any(k.casefold() in sensitive or k.casefold().startswith('x-amz-') for k, _ in parse_qsl(u.query)):
        raise ValueError('Private or signed reference URL is not a shareable citation')
    return value


def timestamp(value):
    try:
        dt = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if dt.tzinfo is None or dt.utcoffset().total_seconds() != 0 or dt > datetime.now(timezone.utc):
            raise ValueError()
    except (TypeError, AttributeError, ValueError):
        raise ValueError('Reference observation requires a real nonfuture UTC timestamp') from None


def selected(data):
    facts = {f['id']: f for f in data['facts']}
    return [facts[m['selected_fact']] for m in data['fields']]


def resolution(value):
    match = re.fullmatch(r'\s*([1-9]\d*)\s*[x×]\s*([1-9]\d*)\s*(?:pixels)?\s*', str(value))
    if not match:
        raise ValueError('Panel resolution requires horizontal x vertical pixels')
    return tuple(map(int, match.groups()))


def panel_check(data):
    facts = {f['attribute']: f for f in selected(data) if f['sheet'] == 'Product Details'}
    active = {name: f for name, f in facts.items() if name in PANEL_FIELDS and f['status'] == 'VERIFIED'}
    if not active:
        return
    native = facts.get('Native Resolution')
    if not native or native['status'] != 'VERIFIED':
        if any(name in RESOLUTION_FIELDS for name in active):
            raise ValueError('Confirmed panel resolution requires verified Native Resolution')
        return
    expected = resolution(native['final_value'])
    for name, fact in active.items():
        if fact['fact_scope'] != 'SOLD_BASE' or fact.get('evidence', {}).get('display_scope') != 'INTERNAL_PANEL':
            raise ValueError('Panel field cannot use external-output/platform capability: ' + name)
        if name in RESOLUTION_FIELDS and resolution(fact['final_value']) != expected:
            raise ValueError('Panel resolution fields disagree with Native Resolution: ' + name)
        if name == 'Display Refresh Rate in Hertz':
            value = fact['final_value']
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                raise ValueError('Internal-panel refresh rate must be a positive numeric Hz value')
        if name == 'Display Resolution Maximum Unit' and str(fact['final_value']).casefold() != 'pixels':
            raise ValueError('Panel resolution unit must be pixels')


def validate(data, audit, ledger_bytes=None, book=None):
    retired_erp = set(json.loads(ERP_POLICY.read_text(encoding='utf-8'))['retiredHosts'])
    if not isinstance(audit, dict) or audit.get('schemaVersion') != 1:
        raise ValueError('New workbook requires reference-audit.json schemaVersion 1')
    for name in ('productId', 'configurationId'):
        if audit.get(name) != data.get(name):
            raise ValueError('Reference audit belongs to another product/configuration')
    if ledger_bytes is not None and audit.get('ledgerSha256') != hashlib.sha256(ledger_bytes).hexdigest():
        raise ValueError('Reference audit is stale for the actual ledger bytes')
    checks = audit.get('references')
    reviews = audit.get('fieldReviews')
    if not isinstance(checks, list) or not isinstance(reviews, list):
        raise ValueError('Reference audit requires browser observations and field reviews')
    references = {}
    for row in checks:
        url = clean_url(row.get('url'))
        if url in references:
            raise ValueError('Duplicate reference observation')
        clean_url(row.get('resolvedUrl'))
        timestamp(row.get('observedAtUtc'))
        if row.get('method') not in {'BROWSER', 'AUTHENTICATED_BROWSER'}:
            raise ValueError('Reference needs an actual browser observation, not a search snippet/HEAD')
        if row.get('outcome') not in {'OPENED', 'NOT_FOUND', 'GONE', 'LOGIN_REQUIRED', 'ACCESS_BLOCKED', 'RENDER_UNAVAILABLE'}:
            raise ValueError('Unknown reference access result')
        if not all(isinstance(row.get(k), str) and row[k].strip() for k in ('title', 'evidence', 'limitations')):
            raise ValueError('Reference needs observed content and access/applicability limitations')
        if row['outcome'] == 'OPENED' and row.get('httpStatus') in {401, 403, 404, 410}:
            raise ValueError('Failed HTTP result cannot be OPENED')
        references[url] = row
    fields = selected(data)
    by_id = {r.get('factId'): r for r in reviews}
    if len(by_id) != len(reviews) or set(by_id) != {f['id'] for f in fields}:
        raise ValueError('Reference field reviews must cover each selected fact exactly once')
    links = {}
    for fact in fields:
        review = by_id[fact['id']]
        if review.get('factSha256') != digest(fact):
            raise ValueError('Reference field review is stale: ' + fact['id'])
        if not all(isinstance(review.get(k), str) and review[k].strip() for k in ('basis', 'supportScope')):
            raise ValueError('Each field needs a substantive applicability/basis review')
        sources = urls(fact['source']) + urls(fact.get('value_candidate_source'))
        # FCC projection appends this exact source to its Source cell.
        if fact['attribute'] == 'FCC ID' and fact.get('evidence', {}).get('fcc_official_url'):
            sources += urls(fact['evidence']['fcc_official_url'])
        sources = list(dict.fromkeys(sources))
        for url in sources:
            clean_url(url)
            if urlsplit(url).hostname in retired_erp:
                raise ValueError('Retired ERP cannot remain an active workbook reference')
            if url not in references or references[url]['outcome'] != 'OPENED':
                raise ValueError('Active reference is unreviewed or cannot open: ' + fact['id'])
        primary = review.get('primaryUrl')
        if sources:
            if primary not in sources or references[primary]['outcome'] != 'OPENED':
                raise ValueError('Field needs an opened primary reference from its actual sources')
        elif primary is not None or review.get('supportScope') not in {'SELLER_POLICY', 'SELLER_RECORD', 'DISPOSITION', 'DERIVED_COPY'}:
            raise ValueError('Non-web field needs an explicit recorded policy/record/disposition basis')
        links[(fact['sheet'], fact['attribute'])] = primary
    panel_check(data)
    if book is not None:
        from candidate_review import data_sheets
        for sheet in data_sheets(book):
            for column in range(2, sheet.max_column + 1):
                key = (sheet.title, sheet.cell(1, column).value)
                for row_index in (5, 7):
                    cell = sheet.cell(row_index, column)
                    active = urls(cell.value)
                    target = links[key] if row_index == 5 else (active[0] if active else None)
                    actual = cell.hyperlink.target if cell.hyperlink else None
                    if active:
                        if target is None or target not in active or actual != target:
                            raise ValueError('Workbook source needs one exact native hyperlink: ' + str(key))
                        clean_url(actual)
                        if actual not in references or references[actual]['outcome'] != 'OPENED':
                            raise ValueError('Workbook hyperlink lacks an opened browser observation')
                    elif actual:
                        raise ValueError('Unexpected hyperlink on a non-web disposition')
                    if cell.data_type == 'f':
                        raise ValueError('Source/candidate reference must remain literal text, not a formula')
    return {'state': 'WORKBOOK_REFERENCE_CONTRACT_PASS', 'referenceCount': len(references),
            'fieldReviewCount': len(fields), 'scope': 'RECORDED_BROWSER_AND_APPLICABILITY_EVIDENCE_NOT_SOURCE_AUTHENTICATION',
            'links': links}


def check_paths(ledger_path, workbook_path=None, required=False):
    ledger_path = Path(ledger_path)
    path = ledger_path.parent/'reference-audit.json'
    if not path.is_file():
        if required:
            raise ValueError('New generation/export/production requires reference-audit.json')
        return {'state': 'LEGACY_REFERENCE_AUDIT_NOT_ASSESSED'}
    raw = ledger_path.read_bytes()
    data = json.loads(raw.decode('utf-8-sig'))
    audit = json.loads(path.read_text(encoding='utf-8-sig'))
    if workbook_path:
        import openpyxl
        book = openpyxl.load_workbook(workbook_path)
        try:
            result = validate(data, audit, raw, book)
        finally:
            book.close()
    else:
        result = validate(data, audit, raw)
    result.pop('links', None)
    return result
